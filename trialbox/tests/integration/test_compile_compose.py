"""Phase 2 DoD against compose: protocol -> /parse -> /compile (>= 20 criteria, CQL≡SQL >= 98 %) -> review.xlsx ->
/approve round trip; plus the all-templates ATOMS gate on HAPI."""

from __future__ import annotations

import io
import os
from datetime import date
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook
from tb_common.config import _read_dotenv
from tb_common.objstore import MinioStore

pytestmark = pytest.mark.integration
ROOT = Path(__file__).resolve().parents[2]
PARSER, COMPILER, LAKE, FHIR = (
    "http://127.0.0.1:8011",
    "http://127.0.0.1:8012",
    "http://127.0.0.1:8013",
    "http://127.0.0.1:8080/fhir",
)


@pytest.fixture(scope="module")
def minio() -> MinioStore:
    env = _read_dotenv(ROOT / "deploy" / ".env")
    return MinioStore("127.0.0.1:9000", env["MINIO_ACCESS_KEY"], env["MINIO_SECRET_KEY"])


def test_compile_and_approve_round_trip(minio: MinioStore) -> None:
    job = "01JITCOMPILE00000000000000"
    pdf = (ROOT / "tests/fixtures/protocols/GZQO_protocol_v3.pdf").read_bytes()
    minio.put(f"attachments/{job}/GZQO_protocol_v3.pdf", pdf)
    parsed = httpx.post(f"{PARSER}/parse", json={"minio_key": f"attachments/{job}/GZQO_protocol_v3.pdf"}, timeout=120)
    assert parsed.status_code == 200, parsed.text
    minio.put(f"attachments/{job}/parsed.json", parsed.content)
    ruleset = f"GZQO-IT{os.getpid() % 1000}"
    resp = httpx.post(
        f"{COMPILER}/compile",
        timeout=1800,
        json={
            "job_id": job,
            "ruleset": ruleset,
            "parsed_doc_key": f"attachments/{job}/parsed.json",
            "kind": "trial",
            "requested_by": "crc1@hospa.test",
            "options": {"index_date": "2026-10-05"},
        },
    )
    assert resp.status_code == 200, resp.text[:2000]
    res = resp.json()
    assert len(res["criteria"]) >= 20
    assert res["tests"]["overall_pct"] >= 98.0 and not res["tests"]["failing"]
    assert res["tests"]["sample_size"] == 200
    wb = load_workbook(io.BytesIO(minio.get(res["review_xlsx_key"])))
    rows = list(wb["review"].iter_rows(min_row=2, values_only=True))
    decisions = [{"id": r[0], "status": r[8]} for r in rows]
    out = httpx.post(
        f"{COMPILER}/approve",
        timeout=1800,
        json={"ruleset": ruleset, "version": res["version"], "by": "crc1@hospa.test", "decisions": decisions},
    )
    assert out.status_code == 200, out.text[:2000]
    body = out.json()
    assert body["status"] == "approved" and body["tag"] == f"{ruleset}/v{res['version']}"


def test_atoms_gate_on_hapi() -> None:
    from criteria_compiler.compile_cql.generator import CqlGenerator, render_common
    from criteria_compiler.compile_cql.translator import Translator, available, fhirhelpers_source
    from criteria_compiler.compile_sql.generator import SqlGenerator
    from criteria_compiler.equivalence.gate import run_gate
    from lake.client import LakeHttp
    from tb_common.fhir import FhirEvaluator, library_resource

    from tests.atoms_ruleset import atoms_ruleset

    if not available():
        pytest.skip("cql-translator jars not available on the host (.cache/jvm)")
    crit, vs, vsi = atoms_ruleset()
    gen = CqlGenerator("ATOMS", "1.0.0", vsi)
    cql = gen.render(crit)
    tr = Translator().translate({"FHIRHelpers": fhirhelpers_source(), "TB_Common": render_common(), gen.library: cql})
    fh = FhirEvaluator(FHIR)
    fh.load(
        [
            library_resource("FHIRHelpers", "4.0.1", fhirhelpers_source(), tr.elm["FHIRHelpers"]),
            library_resource("TB_Common", "1.0.0", render_common(), tr.elm["TB_Common"]),
            library_resource(gen.library, gen.library_version, cql, tr.elm[gen.library]),
        ],
        vs.values(),
    )
    rep = run_gate(
        crit,
        SqlGenerator("ATOMS", "1.0.0", vsi).render(crit),
        gen.library,
        LakeHttp(LAKE),
        fh,
        date(2026, 10, 5),
        "ATOMS|1.0.0",
    )
    assert rep.failing == [], [d.model_dump() for d in (rep.disagreements or [])[:5]]
    assert rep.overall_pct >= 98.0
