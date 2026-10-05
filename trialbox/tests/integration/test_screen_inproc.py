"""Phase 4 in process (CPU CI): SCREEN by email -> encrypted candidate workbook + referral + summary; planted-truth
oracle for structured and note verdicts; MICROBATCH changes sheet on seeded changes; FEEDBACK -> CALIBRATION -> FEAS
uses calibrated rates (SPEC §8.2)."""

from __future__ import annotations

import csv
import io
import json
import shutil
from datetime import date
from pathlib import Path
from typing import Any

import pytest
import tb_contracts as c
from openpyxl import load_workbook
from tb_common.audit import verify
from tb_common.crypto import pid_for_mrn
from tb_common.ruleset import Ruleset

from tests.inproc_box import Box, ingest_site, make_box, unzip_phi

ROOT = Path(__file__).resolve().parents[2]
KEY = b"trialbox-test-site-key-0123456789abcdef"
RUN = date(2026, 10, 5)  # a Monday, the synthetic reference date


@pytest.fixture(scope="module")
def world(synth_dir: Path, tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    tmp = tmp_path_factory.mktemp("screen")
    sec = tmp / "sec"
    sec.mkdir()
    (sec / "site_hmac.key").write_bytes(KEY)
    ingest_site(synth_dir / "site-a", tmp / "lake", sec, "2026-10-04")
    box = make_box(tmp / "box", tmp / "lake", sec, RUN)
    mrns = [r["mrn"] for r in csv.DictReader((synth_dir / "site-a" / "patient.csv").open(encoding="utf-8"))]
    return {"tmp": tmp, "sec": sec, "box": box, "mrn": {pid_for_mrn(KEY, m): m for m in mrns}}


def _list_mail(box: Box, prefix: str, to: str = "pi@hospa.test") -> dict[str, bytes]:
    mail = next(s for s in reversed(box.sent) if s.subject.startswith(prefix) and to in s.to)
    return unzip_phi(box, mail)


def test_screen_by_email(world: dict[str, Any], synth_dir: Path) -> None:
    box: Box = world["box"]
    assert box.mail("SCREEN GZQO version=1.0.0") == "Processed"
    box.orch.drain()
    job = box.orch.db.find(type_="SCREEN")[0]
    assert job.state == "done", job.error
    assert (job.metrics.patients_scoped or 0) > 50 and (job.metrics.llm_calls or 0) > 0  # type: ignore[union-attr]
    files = _list_mail(box, "Candidates GZQO v1.0.0")
    xlsx = next(v for k, v in files.items() if k.endswith(".xlsx"))
    cl = c.CandidateList.model_validate_json(next(v for k, v in files.items() if k.endswith(".json")))
    assert (
        cl.rows
        and {r.tier for r in cl.rows} <= {"high", "review"}
        and cl.summary.high + cl.summary.review == len(cl.rows)
    )
    nexts = [r.next_appointment for r in cl.rows if r.next_appointment]
    assert nexts == sorted(nexts)  # appointment order
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["candidates", "criteria", "feedback"]
    rows = list(wb["candidates"].iter_rows(min_row=2, values_only=True))
    listed = [r for r in rows if r[1]]
    assert all(str(r[0]).isdigit() and len(str(r[0])) == 8 for r in listed[: len(cl.rows)])  # MRN resolved in-box
    assert box.find(f"Done {job.job_id}") and not any(
        s.to == ["sponsor@pharma.example"] for s in box.sent
    )  # summary is the only non-list mail
    # planted truth: structured verdicts == recorded CQL verdicts; note verdicts vs note truth
    expected = json.loads((ROOT / "rulesets/GZQO/tests/expected.json").read_text())["patients"]
    truth = json.loads((synth_dir / "site-a" / "note_truth.json").read_text())
    agree = total = n_agree = n_total = 0
    rs = Ruleset.load(ROOT / "rulesets" / "GZQO")
    kinds = {cr.id: cr.kind for cr in rs.criteria}
    for r in cl.rows:
        mrn = world["mrn"][r.pid]
        for v in r.criteria:
            if mrn in expected and v.id in expected[mrn]:
                pred = expected[mrn][v.id]
                want = "unknown" if pred is None else ("pass" if pred != (kinds[v.id] == "exclusion") else "fail")
                total += 1
                agree += int(v.verdict == want)
            gold = truth.get(mrn, {}).get(v.id)
            if gold is not None and v.verdict != "pending_human":
                truth_pred = {"yes": True, "no": False}.get(gold)
                exp = (
                    "unknown"
                    if truth_pred is None
                    else ("pass" if truth_pred != (kinds[v.id] == "exclusion") else "fail")
                )
                n_total += 1
                n_agree += int(v.verdict == exp)
    assert total > 50 and agree / total >= 0.98, (agree, total)
    assert n_total >= 10 and n_agree / n_total >= 0.90, (n_agree, n_total)  # 300-item set: tools/judge_eval.py
    pool = box.orch.db.pool("GZQO")
    assert {p["pid"] for p in pool} == {r.pid for r in cl.rows}
    referral = [s for s in box.sent if s.subject.startswith("Referral META GZQO")]
    if any(r.department == "META" for r in cl.rows):
        assert referral and referral[0].to == ["meta-head@hospa.test"]
    assert verify(box.audit_dir).ok


def test_microbatch_changes_on_seeded_changes(world: dict[str, Any], synth_dir: Path) -> None:
    box: Box = world["box"]
    pool = box.orch.db.pool("GZQO")
    assert pool, "run after test_screen_by_email"
    by_v = {p["pid"]: {v["id"]: v["verdict"] for v in p["verdicts"]} for p in pool}
    ok = sorted(pid for pid, v in by_v.items() if v.get("GZQO-EXC-04") == "pass" and v.get("GZQO-EXC-11") != "fail")
    assert len(ok) >= 4, by_v
    glp, flare = ok[:2], ok[2:4]
    src = world["tmp"] / "site-a-seeded"
    shutil.copytree(synth_dir / "site-a", src, dirs_exist_ok=True)
    mrn = world["mrn"]
    with (src / "medication.csv").open("a", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        for i, pid in enumerate(glp):
            w.writerow(
                [
                    f"MSEED{i}",
                    mrn[pid],
                    "KC00005209",
                    "A10BJ06",
                    "Semaglutide 1mg/pen",
                    "2026-09-28",
                    "2026-09-28",
                    "2026-10-25",
                    "1",
                    "mg",
                    "active",
                    "META",
                ]
            )
    with (src / "appointment.csv").open("a", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)  # the seeded patients are booked this week
        for i, pid in enumerate(glp + flare):
            w.writerow([f"ASEED{i}", mrn[pid], "2026-10-08 10:00:00", "15", "P12345", "RHEU", "B"])
    with (src / "note.csv").open("a", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        for i, pid in enumerate(flare):
            w.writerow(
                [
                    f"NSEED{i}",
                    mrn[pid],
                    "",
                    "2026-10-03 10:00:00",
                    "progress",
                    "病人今日回診。目前右側第一蹠趾關節急性紅腫熱痛，診斷為急性痛風發作中。給予止痛藥物。",
                ]
            )
    ingest_site(src, world["tmp"] / "lake", world["sec"], "2026-10-05")
    job = box.orch.create(
        c.JobCreate(type="MICROBATCH", ruleset="GZQO", requested_by="scheduler", options={"incident": "0"})
    )
    box.orch.drain()
    job = box.orch.db.get(job.job_id)  # type: ignore[assignment]
    assert job is not None and job.state == "done", job.error
    files = _list_mail(box, "Candidates GZQO v1.0.0")
    xlsx = next(v for k, v in files.items() if k.startswith("this_week_visit1"))
    wb = load_workbook(io.BytesIO(xlsx))
    assert "changes" in wb.sheetnames
    got = {(r[1], r[2], r[3], r[4]) for r in wb["changes"].iter_rows(min_row=2, values_only=True) if r[1]}
    want = {(p, "GZQO-EXC-04", "pass", "fail") for p in glp} | {
        (p, "GZQO-EXC-11", by_v[p]["GZQO-EXC-11"], "fail") for p in flare
    }
    assert got == want
    tiers = {r[1]: r[2] for r in wb["candidates"].iter_rows(min_row=2, values_only=True) if r[1]}
    assert all(tiers[p] == "excluded" for p in glp + flare)


def test_feedback_and_calibration(world: dict[str, Any]) -> None:
    box: Box = world["box"]
    pool = [p["pid"] for p in box.orch.db.pool("GZQO")]
    outcomes = ["enrolled"] * 3 + ["screen_fail"] * 4 + ["declined"] * 5 + ["not_contacted"] * 4
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["pid", "outcome", "reason_code", "note"])
    for pid, o in zip(pool, outcomes, strict=False):
        w.writerow([pid, o, "GZQO-INC-05" if o == "screen_fail" else "", ""])
    w.writerow(["not-a-pid", "enrolled", "", ""])
    box.mail("FEEDBACK GZQO", (("screen_feedback.csv", buf.getvalue().encode()),))
    box.orch.drain()
    fb = box.orch.db.feedback("GZQO")
    assert len(fb) == min(len(pool), len(outcomes))
    done = box.find(f"Done {box.orch.db.find(type_='FEEDBACK')[0].job_id}")[0]
    assert "pid not in the GZQO candidate pool" in done.text()
    box.orch.create(c.JobCreate(type="CALIBRATION", ruleset="GZQO", requested_by="scheduler"))
    box.orch.drain()
    cal = json.loads((box.orch.data_dir / "calibration" / "GZQO.json").read_text())
    assert cal["with_feedback"] == len(fb) and cal["contacted"] == sum(f["outcome"] != "not_contacted" for f in fb)
    expect_calibrated = cal["contacted"] >= cal["min_contacted"]
    assert cal["used_by_feas"] is expect_calibrated
    box.mail("FEAS GZQO")
    box.orch.drain()
    feas = box.orch.db.find(type_="FEAS")[0]
    res = c.FeasibilityResult.model_validate_json(
        box.store.get(
            next(o.minio_key for o in feas.outputs or [] if o.filename.endswith(".json") and ".raw" not in o.filename)
        )
    )
    assert res.simulation.source == ("calibrated" if expect_calibrated else "default")
