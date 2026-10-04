"""Email in -> FEAS -> email out, in process (CPU CI): real mail-gateway intake/outbound, orchestrator, FEAS scenario,
compiler approval loop (fake CQL engine) and the lake loaded with synthetic site A (SPEC §5, §8.1)."""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pytest
import tb_contracts as c
from openpyxl import load_workbook
from pypdf import PdfReader
from tb_common.audit import verify

from tests.inproc_box import Box, make_box

ROOT = Path(__file__).resolve().parents[2]
PDF = ROOT / "tests/fixtures/protocols/GZQO_protocol_v3.pdf"
AUTH = "mx.hospa.test; spf=pass smtp.mailfrom=hospa.test; dkim=pass header.d=hospa.test; dmarc=pass"


@pytest.fixture
def box(ingested: dict[str, Any], site_key_dir: Path, tmp_path: Path) -> Box:
    return make_box(tmp_path, ingested["lake_dir"], site_key_dir)


def test_feas_email_round_trip(box: Box) -> None:
    assert box.mail("FEAS GZQO variant=bmi:25 -- sponsor question") == "Processed"
    queued = box.orch.db.find(type_="FEAS")[0]
    assert box.find(f"Received {queued.job_id}")[0].to == ["crc1@hospa.test"]
    box.orch.drain()
    got = box.orch.db.get(queued.job_id)
    assert got is not None
    assert got.state == "done", got.error
    job = got
    done = box.find(f"Done {job.job_id}")
    assert len(done) == 1 and done[0].to == ["crc1@hospa.test"]
    atts = done[0].attachments()
    names = sorted(atts)
    assert [n.rsplit(".", 1)[-1] for n in names] == ["json", "pdf", "xlsx"]
    pdf = next(v for k, v in atts.items() if k.endswith(".pdf"))
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Eligibility funnel" in text and f"job {job.job_id}" in text and "示範醫院A" in text
    assert b"/FontFile2" in pdf  # CJK font embedded
    result = c.FeasibilityResult.model_validate_json(next(v for k, v in atts.items() if k.endswith(".json")))
    assert result.start_n == 600 and result.funnel[0].criterion_id == "GZQO-INC-01"
    variants = {s.variant for s in result.sensitivity}
    assert variants == {"bmi=25", "unknown_as_pass"}  # the subject option replaces the manifest's bmi variants
    wb = load_workbook(io.BytesIO(next(v for k, v in atts.items() if k.endswith(".xlsx"))))
    assert wb.sheetnames == ["funnel", "sensitivity", "monthly", "criteria", "assumptions"]
    # every output hashed + audited; the unsuppressed raw counts are an output that is never mailed
    assert len(job.outputs or []) == 4 and {o.tag for o in job.outputs or []} == {"aggregate", "phi"}
    assert not any(n.endswith(".raw.json") for n in names)
    assert verify(box.audit_dir).ok
    lines = [json.loads(x) for f in sorted(box.audit_dir.glob("*.jsonl")) for x in f.read_text().splitlines()]
    shas = {x.get("output_sha") for x in lines if x["event"] == "output"}
    assert {o.sha256 for o in job.outputs or []} <= shas
    states = [x["detail"]["to"] for x in lines if x["event"] == "job.state" and x.get("job_id") == job.job_id]
    assert states == ["received", "running", "reporting", "done"]
    # same inputs + versions -> identical machine-readable result
    box.mail("FEAS GZQO variant=bmi:25")
    box.mail("FEAS GZQO")
    box.orch.drain()
    third, second = box.orch.db.find(type_="FEAS")[:2]
    default = c.FeasibilityResult.model_validate_json(
        box.store.get(
            next(o.minio_key for o in third.outputs or [] if o.filename.endswith(".json") and ".raw" not in o.filename)
        )
    )
    assert {s.variant for s in default.sensitivity} == {"bmi=24", "bmi=25", "bmi=27", "unknown_as_pass"}

    def kinds(j: c.Job) -> dict[str, str]:
        return {
            ("raw" if o.filename.endswith(".raw.json") else o.filename.rsplit(".", 1)[1]): o.sha256
            for o in j.outputs or []
        }

    assert kinds(second)["json"] == kinds(job)["json"] and kinds(second)["raw"] == kinds(job)["raw"]
    assert kinds(second)["pdf"] != kinds(job)["pdf"]  # the footer names the job


def test_compile_approve_loop_by_email(box: Box) -> None:
    box.mail("FEAS GZQO-IT", (("GZQO_protocol_v3.pdf", PDF.read_bytes()),))
    box.orch.drain()
    feas = box.orch.db.find(type_="FEAS")[0]
    assert feas.state == "awaiting_approval" and feas.ruleset_version == "1.0.0"
    review = box.find("APPROVE GZQO-IT version=1.0.0 -- review round 1")
    assert len(review) == 1
    atts = review[0].attachments()
    assert sorted(n.rsplit(".", 1)[-1] for n in atts) == ["html", "xlsx", "zip"]
    assert "GZQO-IT-INC-01" in (review[0].msg.get_body(("html",)).get_content())  # type: ignore[union-attr]
    xlsx = next(v for k, v in atts.items() if k.endswith(".xlsx"))
    # reply-to-thread with a free-text subject: the In-Reply-To header identifies the review
    assert box.mail("looks fine to me", (("review.xlsx", xlsx),), in_reply_to=str(review[0].msg["Message-ID"])) == (
        "Processed"
    )
    box.orch.drain()
    appr = box.orch.db.find(type_="APPROVE")[0]
    assert appr.state == "done", appr.error
    assert box.find("Ruleset GZQO-IT v1.0.0 approved")
    feas = box.orch.db.get(feas.job_id)  # type: ignore[assignment]
    assert feas is not None and feas.state == "done", feas.error
    done = box.find(f"Done {feas.job_id}")
    assert done and any(n.endswith(".pdf") for n in done[0].attachments())


def test_rejections_status_cancel(box: Box) -> None:
    assert box.mail("FEAS GZQO", sender="stranger@evil.example", auth="mx.hospa.test; spf=fail") == "Rejected"
    assert not box.sent  # no backscatter to unauthenticated unknown senders
    assert box.mail("FEAS GZQO", auth="mx.hospa.test; spf=fail; dkim=fail") == "Rejected"
    assert "authentication failed" in box.sent[-1].text()
    assert box.mail("PLEASE COUNT GZQO") == "Rejected"
    assert "unknown command" in box.sent[-1].text()
    assert box.mail("APPROVE GZQO version=1.0.0", sender="nurse-rheu@hospa.test") == "Rejected"
    assert "not permitted" in box.sent[-1].text()
    box.mail("FEAS NOPE")
    job = box.orch.db.find(type_="FEAS")[0]
    box.orch.drain()
    assert box.orch.db.get(job.job_id).state == "failed"  # type: ignore[union-attr]
    box.mail(f"STATUS {job.job_id}")
    failed = box.find(f"Failed {job.job_id}")[0]
    assert "Step: **ruleset**" in failed.text() or "Step: ruleset" in failed.text()
    status = box.find(f"Status {job.job_id}")
    assert status and '"job_id"' in status[0].text()
    box.mail("FEAS GZQO")
    queued = box.orch.db.find(type_="FEAS")[0]
    box.mail(f"CANCEL {queued.job_id}")  # STATUS / CANCEL are answered at once, ahead of the queue
    box.orch.drain()
    got = box.orch.db.get(queued.job_id)
    assert got is not None and got.state == "failed" and got.error is not None and got.error.step == "cancelled"
