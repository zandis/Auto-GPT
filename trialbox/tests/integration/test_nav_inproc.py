"""Phase 5 in process (CPU CI): ``NAV RA-BIO dept=RHEU`` by email -> four lists, application drafts with [待補]
markers, NavLists JSON; list semantics checked against the planted claim/approval states (SPEC §8.4, §9.3, §9.4)."""

from __future__ import annotations

import csv
import io
import json
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pytest
import tb_contracts as c
from openpyxl import load_workbook
from tb_common.crypto import pid_for_mrn

from tests.inproc_box import Box, ingest_site, make_box, unzip_phi

KEY = b"trialbox-test-site-key-0123456789abcdef"
RUN = date(2026, 10, 5)


@pytest.fixture(scope="module")
def world(synth_dir: Path, tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    tmp = tmp_path_factory.mktemp("nav")
    sec = tmp / "sec"
    sec.mkdir()
    (sec / "site_hmac.key").write_bytes(KEY)
    ingest_site(synth_dir / "site-a", tmp / "lake", sec, "2026-10-04")
    box = make_box(tmp / "box", tmp / "lake", sec, RUN)
    mrns = [r["mrn"] for r in csv.DictReader((synth_dir / "site-a" / "patient.csv").open(encoding="utf-8"))]
    return {"box": box, "mrn": {pid_for_mrn(KEY, m): m for m in mrns}}


def test_nav_by_email(world: dict[str, Any], synth_dir: Path) -> None:
    import docx

    box: Box = world["box"]
    assert box.mail("NAV RA-BIO dept=RHEU", sender="nurse-rheu@hospa.test") == "Processed"
    box.orch.drain()
    job = box.orch.db.find(type_="NAV")[0]
    assert job.state == "done", job.error
    mail = next(s for s in box.sent if s.subject.startswith("NAV RA-BIO RHEU"))
    assert sorted(mail.to) == ["nurse-rheu@hospa.test", "rheu-dr@hospa.test"]
    files = unzip_phi(box, mail)
    wb = load_workbook(io.BytesIO(next(v for k, v in files.items() if k.endswith(".xlsx"))))
    assert wb.sheetnames == ["likely_eligible", "renewal_due", "doc_gaps", "maybe_ineligible"]
    nav = c.NavLists.model_validate_json(next(v for k, v in files.items() if k.endswith(".json")))
    lists = nav.lists
    assert lists.likely_eligible and lists.renewal_due, {k: len(v) for k, v in lists.model_dump().items()}
    states = json.loads((synth_dir / "site-a" / "states.json").read_text())
    mrn = world["mrn"]
    for r in lists.likely_eligible:
        st = states[mrn[r.pid]]
        assert st.get("approval") in ("none", "expired"), st  # no active approval, no recent application
        v = {x.id: x.verdict for x in r.criteria}
        assert all(v[f"RA-BIO-INC-0{i}"] == "pass" for i in range(1, 7))
        assert not any(v[k] == "fail" for k in v if "-EXC-" in k)
    for r in lists.renewal_due:
        st = states[mrn[r.pid]]
        assert st.get("approval") == "renewal_due" and r.approval_end is not None
        assert r.approval_end <= RUN + timedelta(days=60)
    for r in lists.maybe_ineligible:
        assert states[mrn[r.pid]].get("approval") in ("active_far", "renewal_due")
    gap_ids = {r.pid for r in lists.doc_gaps}
    for r in lists.likely_eligible + lists.renewal_due:
        has_gap = any(m.criterion_id and "-DOC-" in m.criterion_id for m in r.missing)
        assert has_gap == (r.pid in gap_ids)
        if has_gap:
            assert all(m.suggested_order for m in r.missing if m.criterion_id and "-DOC-" in m.criterion_id)
    drafts = {k: v for k, v in files.items() if k.endswith(".docx")}
    assert len(drafts) == len(lists.likely_eligible) + len(lists.renewal_due)
    for data in drafts.values():  # opens as a Word document; unknown fields highlighted [待補]
        d = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(
            cell.text for t in d.tables for row in t.rows for cell in row.cells
        )
        assert "附表十五" in text and "類風濕" in text and "{{" not in text and "{%" not in text
        assert "DAS28" in text
        marks = [r for p in d.paragraphs for r in p.runs if r.text == "[待補]"] + [
            r
            for t in d.tables
            for row in t.rows
            for cell in row.cells
            for p in cell.paragraphs
            for r in p.runs
            if r.text == "[待補]"
        ]
        assert marks and all("yellow" in r._r.xml for r in marks)  # shaded yellow, bold
    summary = next(s for s in box.sent if s.subject.startswith(f"Done {job.job_id}"))
    assert any(k.endswith(".pdf") for k in summary.attachments())
