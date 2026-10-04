"""Phase 7 in process (CPU CI): ``COHORT GOUT-COH`` on two synthetic hospitals; the member box (B) mails its alliance
table to the root (A) as ``COHORT MERGE``; A returns the merged table. Counts are checked against the raw fixture
exports (diagnoses, consent registry); the trial simulation covers the three recruiting gout trials of the
ClinicalTrials.gov cassette (SPEC §8.3)."""

from __future__ import annotations

import csv
import io
from datetime import date
from email import policy
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from orchestrator.scenarios import cohort
from pypdf import PdfReader

from tests.inproc_box import SITE_B, Box, Sent, ingest_site, make_box

RUN = date(2026, 10, 5)
KEYS = {"a": b"trialbox-test-site-key-0123456789abcdef", "b": b"trialbox-test-site-key-B-0123456789abcd"}
AUTH_B = "mx.hospb.test; spf=pass smtp.mailfrom=hospb.test; dkim=pass header.d=hospb.test; dmarc=pass"
AUTH_A_FROM_B = "mx.hospa.test; spf=pass smtp.mailfrom=hospb.test; dkim=pass header.d=hospb.test; dmarc=pass"


@pytest.fixture(scope="module")
def boxes(synth_dir: Path, tmp_path_factory: pytest.TempPathFactory) -> dict[str, Box]:
    tmp = tmp_path_factory.mktemp("cohort")
    out: dict[str, Box] = {}
    for site in ("a", "b"):
        sec = tmp / f"sec-{site}"
        sec.mkdir()
        (sec / "site_hmac.key").write_bytes(KEYS[site])
        ingest_site(synth_dir / f"site-{site}", tmp / f"lake-{site}", sec, "2026-10-04", f"DEMO-{site.upper()}")
        out[site] = make_box(
            tmp / f"box-{site}", tmp / f"lake-{site}", sec, RUN, settings_update=SITE_B if site == "b" else None
        )
    return out


def _csv_rows(data: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(data.decode("utf-8"))))


def _oracle(src: Path, q_end: date) -> tuple[set[str], set[str]]:
    """Gout population (any M10/M1A diagnosis dated on or before the quarter end) and consented MRNs."""
    pop = {
        r["mrn"]
        for r in csv.DictReader((src / "diagnosis.csv").open(encoding="utf-8"))
        if r["icd10"].startswith(("M10", "M1A")) and date.fromisoformat(r["diag_date"][:10]) <= q_end
    }
    consent = {
        r["mrn"]
        for r in csv.DictReader((src / "registry.csv").open(encoding="utf-8"))
        if r["consent_contact"] == "Y" and date.fromisoformat(r["consent_date"]) <= q_end
    }
    return pop, consent


def _done(box: Box, subject: str, job_type: str = "COHORT", **kw: Any) -> tuple[Any, Sent]:
    assert box.mail(subject, **kw) == "Processed"
    box.orch.drain()
    job = box.orch.db.find(type_=job_type)[0]
    assert job.state == "done", job.error
    mail = next(s for s in reversed(box.sent) if s.subject.startswith(f"Done {job.job_id}") and s.attachments())
    return job, mail


def test_cohort_site_a(boxes: dict[str, Box], synth_dir: Path) -> None:
    box = boxes["a"]
    _, mail = _done(box, "COHORT GOUT-COH")
    assert sorted(mail.to) == ["crc1@hospa.test", "pi@hospa.test"]
    files = mail.attachments()
    table = _csv_rows(next(v for k, v in files.items() if k.startswith("cohort_table_GOUT-COH_DEMO-A_2026Q3")))
    assert list(table[0]) == list(cohort.COLUMNS)
    quarters = sorted({r["quarter"] for r in table})
    assert quarters == ["2025Q4", "2026Q1", "2026Q2", "2026Q3"]
    for q in quarters:
        q_end = cohort.completed_quarters(RUN, 4)[quarters.index(q)]
        pop, consent = _oracle(synth_dir / "site-a", q_end)
        rows = {r["criterion_id"]: r for r in table if r["quarter"] == q}
        assert int(rows["GOUT-COH-INC-01"]["n"]) == len(pop)
        assert int(rows["GOUT-COH-INC-01"]["n_contactable"]) == len(pop & consent)
        assert set(rows) == {f"GOUT-COH-INC-0{i}" for i in range(1, 8)} | {
            "GOUT-COH-INC-02+GOUT-COH-INC-03",
            "GOUT-COH-INC-04+GOUT-COH-INC-05",
        }
        for r in rows.values():  # every count is a non-negative integer or a suppressed '<5'
            for col in ("n", "n_contactable"):
                v = r[col]
                assert v == "<5" or (v.isdigit() and (int(v) == 0 or int(v) >= 5)), (r["criterion_id"], col, v)
                if v.isdigit():
                    assert int(v) <= int(rows["GOUT-COH-INC-01"][col])
        combo = rows["GOUT-COH-INC-02+GOUT-COH-INC-03"]["n"]
        if combo.isdigit() and rows["GOUT-COH-INC-02"]["n"].isdigit():
            assert int(combo) <= int(rows["GOUT-COH-INC-02"]["n"])
    pdf = next(v for k, v in files.items() if k.endswith(".pdf"))
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Cohort report" in text and "NCT99000001" in text and "2026Q3" in text
    wb = load_workbook(io.BytesIO(next(v for k, v in files.items() if k.startswith("trial_sim_"))))
    trials = list(wb["trials"].iter_rows(values_only=True))
    assert [r[0] for r in trials[1:]] == ["NCT99000001", "NCT99000002", "NCT99000003"]
    for values in trials[1:]:
        row: dict[str, Any] = dict(zip([str(h) for h in trials[0]], values, strict=True))
        assert row["criteria_total"] >= 7 and 0 < row["criteria_counted"] <= row["criteria_total"]
        assert row["cached"] == "no"
        if isinstance(row["eligible_now"], int):
            assert row["enrol_12m_P10"] <= row["enrol_12m_P50"] <= row["enrol_12m_P90"]
    steps = list(wb["steps"].iter_rows(values_only=True))
    assert any(s[0] == "NCT99000001" and s[3] == "no" for s in steps[1:])  # the human criterion is not counted
    # second run: compiles come from the cache (NCT id + last update)
    _, mail2 = _done(box, "COHORT GOUT-COH lookback=1")
    wb2 = load_workbook(io.BytesIO(next(v for k, v in mail2.attachments().items() if k.startswith("trial_sim_"))))
    assert {r[16] for r in list(wb2["trials"].iter_rows(values_only=True))[1:]} == {"yes"}


def test_member_shares_and_root_merges(boxes: dict[str, Box], synth_dir: Path) -> None:
    a, b = boxes["a"], boxes["b"]
    b.mail("COHORT GOUT-COH lookback=1", sender="crc1@hospb.test", auth=AUTH_B)
    b.orch.drain()
    bjob = b.orch.db.find(type_="COHORT")[0]
    assert bjob.state == "done", bjob.error
    share = next(s for s in b.sent if s.subject.startswith("COHORT MERGE"))
    assert share.to == ["trialbox@hospa.test"]
    shared = share.attachments()
    assert list(shared) == ["cohort_table_GOUT-COH_DEMO-B_2026Q3.csv"]
    b_rows = _csv_rows(shared["cohort_table_GOUT-COH_DEMO-B_2026Q3.csv"])
    pop_b, _ = _oracle(synth_dir / "site-b", date(2026, 9, 30))
    assert int(next(r for r in b_rows if r["criterion_id"] == "GOUT-COH-INC-01")["n"]) == len(pop_b)
    # the root's MTA adds Authentication-Results; the mail arrives at A's intake mailbox
    msg = share.msg
    del msg["Authentication-Results"]
    msg["Authentication-Results"] = AUTH_A_FROM_B
    assert a.gw.handle(msg.as_bytes(policy=policy.SMTP)) == "Processed"
    a.orch.drain()
    merge = a.orch.db.find(type_="COHORT", ruleset="MERGE")[0]
    assert merge.state == "done", merge.error
    out = next(s for s in reversed(a.sent) if s.subject.startswith(f"Done {merge.job_id}") and s.attachments())
    assert "trialbox@hospb.test" in out.to
    merged = _csv_rows(next(iter(out.attachments().values())))
    pop = {r["site_id"]: r for r in merged if r["criterion_id"] == "GOUT-COH-INC-01" and r["quarter"] == "2026Q3"}
    assert set(pop) == {"DEMO-A", "DEMO-B", "ALLIANCE"}
    pop_a, _ = _oracle(synth_dir / "site-a", date(2026, 9, 30))
    assert int(pop["ALLIANCE"]["n"]) == len(pop_a) + len(pop_b)
    assert int(pop["DEMO-A"]["n"]) == len(pop_a) and int(pop["DEMO-B"]["n"]) == len(pop_b)
    assert {r["definition_version"] for r in merged} == {"GOUT-COH@1.0.0"}


def test_merge_rejects_bad_tables_and_non_root(boxes: dict[str, Box]) -> None:
    a, b = boxes["a"], boxes["b"]
    bad = b"site_id,disease,quarter\nDEMO-X,gout,2026Q3\n"
    a.mail("COHORT MERGE", attachments=(("cohort_table_bad.csv", bad),), sender="crc1@hospa.test")
    a.orch.drain()
    job = a.orch.db.find(type_="COHORT", ruleset="MERGE")[0]
    assert job.state == "failed" and job.error is not None and "header must be" in job.error.message
    b.mail("COHORT MERGE", attachments=(("t.csv", bad),), sender="crc1@hospb.test", auth=AUTH_B)
    b.orch.drain()
    job = b.orch.db.find(type_="COHORT", ruleset="MERGE")[0]
    assert job.state == "failed" and job.error is not None and "not the alliance root" in job.error.message
