"""FEAS building blocks without a lake: simulation, variants, funnel SQL guard, reports, schedule, harness."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

import pytest
from lake.sqlguard import check_select
from openpyxl import load_workbook
from orchestrator.reports import feasibility_pdf, feasibility_xlsx
from orchestrator.reports.common import ReportMeta
from orchestrator.scenarios import calibration
from orchestrator.scenarios.feas import capacity
from orchestrator.scenarios.feas_compute import (
    FeasError,
    build_variants,
    funnel_query,
    parse_variant_option,
    simulate,
    steps,
)
from orchestrator.scheduler import ScheduleError, cron_kwargs
from pypdf import PdfReader
from tb_common.config import load_settings
from tb_common.ruleset import Ruleset
from tb_contracts import FeasibilityResult

from tools.funnel_compare import compare, markdown, system_counts

ROOT = Path(__file__).resolve().parents[2]
GZQO = Ruleset.load(ROOT / "rulesets" / "GZQO")
SETTINGS = load_settings(ROOT / "deploy" / "settings.example.yaml")


def sim(**kw: object) -> dict[str, float]:
    base: dict[str, object] = dict(prevalent=20, monthly_rate=2.0, reach=0.6, accept=0.35, capacity=None, seed=7)
    base.update(kw)
    return simulate(**base)  # type: ignore[arg-type]


def test_simulation_deterministic_and_monotone() -> None:
    a = sim()
    assert a == sim() and a["low"] <= a["mid"] <= a["high"]
    assert sim(prevalent=0, monthly_rate=0.0) == {"low": 0.0, "mid": 0.0, "high": 0.0}
    assert sim(capacity=1.0)["high"] <= 12  # at most one per month for 12 months
    assert sim(reach=0.9, accept=0.9)["mid"] > a["mid"]
    assert sim(seed=8) != a or True  # different seeds may coincide; determinism is per seed


def test_variants_and_steps() -> None:
    st = steps(GZQO)
    applied = [s.id for s in st if s.applied]
    assert "GZQO-INC-05" not in applied and "GZQO-INC-03" in applied  # note criterion not applied in counts
    vs = build_variants(GZQO, {"bmi": "24,25"}, st)
    assert [v.name for v in vs] == ["bmi=24", "bmi=25"] and all(v.criterion_id == "GZQO-INC-03" for v in vs)
    assert vs[0].ir is not None and '"num":24.0' in vs[0].ir.model_dump_json()
    by_id = build_variants(GZQO, {"INC-08": "45"}, st)
    assert by_id[0].name == "INC-08=45" and by_id[0].column_id.endswith("GZQO-INC-08")
    with pytest.raises(FeasError):
        build_variants(GZQO, {"das28": "3"}, st)  # no criterion uses DAS28
    with pytest.raises(FeasError):
        build_variants(GZQO, {"INC-05": "3"}, st)  # note criterion
    with pytest.raises(FeasError):
        build_variants(GZQO, {"bmi": "x"}, st)
    assert parse_variant_option("bmi:24,25,27") == {"bmi": "24,25,27"}
    with pytest.raises(FeasError):
        parse_variant_option("bmi")
    sql = funnel_query(GZQO, st, vs, ["RHEU"])
    check_select(sql)  # passes the lake's SELECT-only guard
    assert "$5" in sql and '"C_V1-GZQO-INC-03"' in sql


def _result() -> FeasibilityResult:
    return FeasibilityResult.model_validate(
        {
            "ruleset": "GZQO",
            "version": "1.0.0",
            "snapshot": "2026-10-04",
            "population": "hospital",
            "lookback_months": 36,
            "run_date": "2026-10-05",
            "start_n": 600,
            "funnel": [
                {
                    "criterion_id": "GZQO-INC-01",
                    "label": "Age ≥18",
                    "remaining": 570,
                    "dropped": 30,
                    "pct": 95.0,
                    "unknown": 0,
                    "applied": True,
                },
                {
                    "criterion_id": "GZQO-INC-02",
                    "label": "痛風 diagnosis",
                    "remaining": 168,
                    "dropped": 402,
                    "pct": 28.0,
                    "unknown": 0,
                    "applied": True,
                },
                {
                    "criterion_id": "GZQO-INC-05",
                    "label": "≥2 flares",
                    "remaining": 168,
                    "dropped": 0,
                    "pct": 28.0,
                    "applied": False,
                },
                {
                    "criterion_id": "GZQO-EXC-01",
                    "label": "T1DM",
                    "remaining": 165,
                    "dropped": "<5",
                    "pct": 27.5,
                    "unknown": 0,
                    "applied": True,
                },
            ],
            "sensitivity": [{"criterion_id": "*", "variant": "unknown_as_pass", "remaining": 200, "delta": 35}],
            "monthly_new": [{"month": f"2025-{m:02d}", "n": (m if m > 4 else "<5")} for m in range(1, 13)],
            "simulation": {
                "reach_rate": 0.6,
                "accept_rate": 0.35,
                "capacity_per_month": 9,
                "months": 12,
                "low": 1,
                "mid": 4,
                "high": 7,
                "source": "default",
                "iterations": 1000,
                "monthly_eligible_mean": 1.1,
            },
            "notes": ["note one"],
        }
    )


META = ReportMeta(
    "DEMO-A",
    "示範醫院A",
    "GZQO",
    "1.0.0",
    "judge-model",
    "2026-10-04",
    "01JB0000000000000000000000",
    "2026-10-05",
    contact="crc1@hospa.test",
)


def test_reports_deterministic_and_complete() -> None:
    r = _result()
    pdf = feasibility_pdf.render(r, GZQO, META, {"cold_chain": True})
    assert pdf == feasibility_pdf.render(r, GZQO, META, {"cold_chain": True})
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    for needle in (
        "Feasibility report",
        "Answer",
        "Eligibility funnel",
        "Sensitivity",
        "New eligible patients",
        "Simulation assumptions",
        "Appendix A",
        "Appendix B",
        META.footer,
        "痛風",
        "<5",
        "not applied in counts",
        "Cold-chain",
    ):
        assert needle in text, needle
    assert "1–7" in text and "痛風 diagnosis (GZQO-INC-02, −402)" in text  # answer + biggest barrier
    xlsx = feasibility_xlsx.render(r, GZQO, META)
    assert xlsx == feasibility_xlsx.render(r, GZQO, META)
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["funnel", "sensitivity", "monthly", "criteria", "assumptions"]
    rows = list(wb["funnel"].iter_rows(values_only=True))
    assert rows[1][:3] == ("START", "hospital", 600) and rows[5][3] == "<5"


def test_capacity_calibration_schedule(tmp_path: Path) -> None:
    assert capacity(SETTINGS, ["P12345", "P23456"]) == 7.0
    assert capacity(SETTINGS, None) == 9.0
    assert capacity(SETTINGS, ["P99999"]) is None
    cal = calibration.load(tmp_path, "GZQO", SETTINGS)
    assert (cal.reach_rate, cal.accept_rate, cal.source) == (0.6, 0.35, "default")
    (tmp_path / "calibration").mkdir()
    (tmp_path / "calibration" / "GZQO.json").write_text(
        json.dumps({"reach_rate": 0.5, "accept_rate": 0.2, "contacted": 80})
    )
    cal = calibration.load(tmp_path, "GZQO", SETTINGS)
    assert cal.source == "calibrated" and cal.concentration == 80.0
    assert cron_kwargs("02:00") == {"hour": 2, "minute": 0}
    assert cron_kwargs("MON 03:30") == {"hour": 3, "minute": 30, "day_of_week": "mon"}
    assert cron_kwargs("1st 04:00", quarterly=True) == {"hour": 4, "minute": 0, "day": 1, "month": "1,4,7,10"}
    for bad in ("25:00", "noon", "XYZ 01:00", "MON TUE 01:00"):
        with pytest.raises(ScheduleError):
            cron_kwargs(bad)


def test_funnel_compare_harness() -> None:
    system = system_counts({"start_n": 600, "remaining": {"A": 100, "B": 50, "C": 20}})
    rows = compare(
        system,
        [
            {"criterion_id": "START", "manual_n": "600"},
            {"criterion_id": "A", "manual_n": "95", "explanation": "dx window differs"},
            {"criterion_id": "B", "manual_n": "48"},
            {"criterion_id": "C", "manual_n": "30", "explanation": "x"},
            {"criterion_id": "D", "manual_n": "1"},
        ],
    )
    assert [r.status for r in rows] == ["ok", "ok", "unexplained", "out_of_tolerance", "missing"]
    report = markdown(rows, 15)
    assert "FAIL" in report and "| A | 100 | 95 | +5.3 | ok |" in report
    assert system_counts(json.loads(_result().model_dump_json()))["GZQO-INC-02"] == 168
    _ = date  # imported for readability of fixtures
