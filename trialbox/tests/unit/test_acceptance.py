"""Acceptance runner (SPEC §11.2; tools/acceptance.py): stratified blinded sampling, unblinded metrics with
stratum weights, criterion agreement, TWPAS / ingest / reproducibility checks and the report."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from tools import acceptance as acc


def _cands(n_high: int, n_review: int, n_excl: int) -> dict[str, Any]:
    rows = []
    for tier, n in (("high", n_high), ("review", n_review), ("excluded", n_excl)):
        for i in range(n):
            rows.append(
                {
                    "pid": f"{tier}-{i:03d}",
                    "tier": tier,
                    "criteria": [
                        {"id": "X-INC-01", "label": "age", "class": "structured", "verdict": "pass"},
                        {"id": "X-INC-02", "label": "note", "class": "note", "verdict": "pass" if i % 2 else "fail"},
                    ],
                }
            )
    return {"rows": rows}


def test_sampling_is_stratified_blinded_and_reproducible() -> None:
    cands = _cands(30, 40, 50)
    scoped = [r["pid"] for r in cands["rows"]] + [f"other-{i}" for i in range(300)]
    pids, key = acc.screen_sample(cands, scoped, 120)
    assert len(pids) >= 100 and len(set(pids)) == len(pids)
    assert all(len(key["strata"][s]["sampled"]) >= 10 for s in acc.STRATA)
    assert key["strata"]["not_listed"]["population"] == 300
    assert pids == acc.screen_sample(cands, scoped, 120)[0]  # seeded
    assert pids != sorted(pids)  # shuffled: no stratum order visible to the CRC


def test_screen_metrics_weighted(tmp_path: Path) -> None:
    cands = _cands(20, 20, 20)
    scoped = [r["pid"] for r in cands["rows"]] + [f"other-{i}" for i in range(100)]
    pids, key = acc.screen_sample(cands, scoped, 100)
    # CRC: every high/review patient eligible except review-000; nobody else eligible; criterion verdicts copy the
    # system except one note disagreement
    adjudication = []
    for p in pids:
        eligible = p.startswith(("high", "review")) and p != "review-000"
        crit = {"X-INC-01": "pass", "X-INC-02": "pass" if int(p.split("-")[-1]) % 2 else "fail"} if "-" in p else {}
        if p == "high-001":
            crit["X-INC-02"] = "unknown"
        adjudication.append({"pid": p, "eligible": eligible, "minutes": 6.0, "criteria": crit})
    res = acc.check_screen(cands, adjudication, key)
    t = res["trial"]
    conf = t["confusion_weighted"]
    assert conf["default"]["fp"] == 0 and conf["review"]["fn"] == 0  # every eligible patient is high or review
    assert "spec 100.0 %" in t["summary"] and "review tier sens 100.0 %" in t["summary"]
    # default tier misses the eligible review patients -> low sensitivity -> trial check fails
    assert not t["passed"]
    assert "structured 100.0 %" in res["criterion"]["summary"]
    listed = [p for p in pids if not p.startswith("other")]
    expected_note = 100.0 * (len(listed) - ("high-001" in listed)) / len(listed)
    assert f"note {round(expected_note, 1)} %" in res["criterion"]["summary"]
    assert res["time"]["passed"]


def test_adjudication_workbook_round_trip(tmp_path: Path) -> None:
    data = acc.write_adjudication(["p1", "p2"], [("X-INC-01", "age")], {"p1": "10000001"})
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    ws = wb["adjudication"]
    ws["C2"] = "Y"
    ws["D2"] = 7
    ws["F2"] = "pass"
    buf = io.BytesIO()
    wb.save(buf)
    rows = acc.read_adjudication(buf.getvalue())
    assert rows[0] == {"pid": "p1", "eligible": True, "minutes": 7.0, "criteria": {"X-INC-01": "pass"}}
    assert rows[1]["eligible"] is None


def test_twpas_ingest_reproduce_and_report(tmp_path: Path) -> None:
    good = {
        "validator": "hl7-validator",
        "bundles": [{"pid": "a", "file": "f", "validator_errors": [], "precheck": {"passed": True}}],
    }
    assert acc.check_twpas([good])["passed"]
    structural = {**good, "validator": "structural"}
    assert not acc.check_twpas([structural])["passed"]
    assert acc.check_ingest({"passed": True, "validation": {"error_pct": 0.1, "validator": "hl7-validator"}})["passed"]
    assert not acc.check_ingest({"passed": True, "validation": {"error_pct": 0.6}})["passed"]
    rep = acc.check_reproduce([{"job_id": "J", "type": "FEAS", "identical": True, "files": [{}, {}]}])
    assert rep["passed"] and "1/1 jobs identical over 2" in rep["summary"]
    out = tmp_path / "acc"
    acc._save(out, "ingest", {"passed": True, "summary": "ok"})
    md = acc.report(out)
    assert "| ingest |" in md and "PASS" in md and "not run" in md and "Overall: NOT YET" in md
    assert json.loads((out / "ingest.json").read_text())["target"].startswith("nightly run")
