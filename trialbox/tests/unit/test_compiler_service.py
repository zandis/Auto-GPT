"""criteria-compiler orchestration with injected fakes (no JVM / HAPI needed): approval loop rules (SPEC §5)."""

from __future__ import annotations

import io
import shutil
from datetime import date
from pathlib import Path
from typing import Any

import pytest
import tb_contracts as c
from criteria_compiler.compile_cql.translator import Translation
from criteria_compiler.repo import RulesetRepo
from criteria_compiler.review import parse_review_xlsx
from criteria_compiler.service import CompileFailed, Compiler, CompilerDeps
from criteria_compiler.terminology.mapper import Terminology
from doc_parser.parser import parse
from fastapi.testclient import TestClient
from llm_stub.app import app as stub
from openpyxl import load_workbook
from tb_common.audit import AuditLog, verify
from tb_common.llm import LlmClient
from tb_common.objstore import FsStore

ROOT = Path(__file__).resolve().parents[2]
PDF = ROOT / "tests/fixtures/protocols/GZQO_protocol_v3.pdf"


class FakeTranslator:
    def translate(self, sources: dict[str, str], check: set[str] | None = None) -> Translation:
        return Translation({name: '{"library":{}}' for name in sources})


class FakeGate:
    def __init__(self) -> None:
        self.failing: list[str] = []
        self.calls = 0

    def __call__(
        self, criteria: list[c.CriterionIR], sql: str, library: str, index_date: date, seed: str
    ) -> c.EquivalenceReport:
        self.calls += 1
        ids = [x.id for x in criteria if x.class_ == "structured"]
        per = {i: (50.0 if i in self.failing else 100.0) for i in ids}
        return c.EquivalenceReport(
            overall_pct=100.0 if not self.failing else 90.0,
            per_criterion=per,
            failing=[i for i in ids if i in self.failing],
            sample_size=200,
            engine="fake",
        )


@pytest.fixture
def env(tmp_path: Path) -> dict[str, Any]:
    seed = tmp_path / "seed"
    shutil.copytree(ROOT / "rulesets", seed, ignore=shutil.ignore_patterns("GZQO", "RA-BIO"))
    store = FsStore(tmp_path / "obj")
    audit = AuditLog(tmp_path / "audit")
    llm = LlmClient("http://stub/v1", "stub", audit=audit)
    llm.http = TestClient(stub)
    gate = FakeGate()
    deps = CompilerDeps(
        store=store,
        repo=RulesetRepo(tmp_path / "repo", seed),
        llm=llm,
        term=Terminology(),
        settings_thresholds={"equivalence_min_pct": 98.0},
        translator=FakeTranslator(),  # type: ignore[arg-type]
        audit=audit,
        equivalence=gate,
        mrn_regex=r"^\d{8}$",
    )
    data = PDF.read_bytes()
    store.put("attachments/J/parsed.json", c.dump_json(parse(data, PDF.name)).encode())
    return {"comp": Compiler(deps), "store": store, "gate": gate, "deps": deps, "tmp": tmp_path}


def _compile(env: dict[str, Any], key: str = "attachments/J/parsed.json", **kw: Any) -> c.CompileResult:
    comp: Compiler = env["comp"]
    return comp.compile(
        c.CompileRequest(
            job_id="01JTEST0000000000000000000",
            ruleset="GZQO",
            parsed_doc_key=key,
            requested_by="crc1@hospa.test",
            options={"index_date": "2026-10-05"},
            **kw,
        )
    )


def _decisions(store: FsStore, key: str | None, changes: dict[str, tuple[str, str | None]] | None = None) -> list[Any]:
    assert key is not None
    wb = load_workbook(io.BytesIO(store.get(key)))
    ws = wb["review"]
    for row in ws.iter_rows(min_row=2):
        cid = str(row[0].value)
        if changes and cid in changes:
            row[8].value, row[9].value = changes[cid]
    buf = io.BytesIO()
    wb.save(buf)
    return parse_review_xlsx(buf.getvalue())[0]


def _approve(env: dict[str, Any], decisions: list[Any]) -> c.ApproveResult:
    comp: Compiler = env["comp"]
    return comp.approve(c.ApproveRequest(ruleset="GZQO", version="1.0.0", by="crc1@hospa.test", decisions=decisions))


def test_compile_produces_draft_and_review(env: dict[str, Any]) -> None:
    res = _compile(env)
    assert len(res.criteria) == 23 and res.version == "1.0.0"
    assert sum(1 for x in res.criteria if x.class_ == "structured") == 18
    assert res.manifest.status == "draft" and res.manifest.scopes.screen is not None
    assert res.manifest.routing.aggregate_to == ["crc1@hospa.test"]
    assert res.branch == "draft/01JTEST0000000000000000000" and res.branch in env["deps"].repo.branches()
    assert res.review_xlsx_key and env["store"].exists(res.review_xlsx_key)
    assert not res.phi_guard_hit and not res.needs_review
    assert verify(env["tmp"] / "audit").ok


def test_approve_all_tags_and_materializes(env: dict[str, Any]) -> None:
    res = _compile(env)
    out = _approve(env, _decisions(env["store"], res.review_xlsx_key))
    assert out.status == "approved" and out.tag == "GZQO/v1.0.0"
    assert out.manifest is not None and out.manifest.status == "approved"
    assert out.manifest.review is not None and out.manifest.review.approved_by == "crc1@hospa.test"
    assert (env["tmp"] / "repo" / "GZQO" / "manifest.yaml").exists()


def test_edit_round_then_approve(env: dict[str, Any]) -> None:
    res = _compile(env)
    r1 = _approve(
        env,
        _decisions(
            env["store"],
            res.review_xlsx_key,
            {"GZQO-INC-03": ("edit", "Body mass index (BMI) ≥30 kg/m² at screening.")},
        ),
    )
    assert r1.status == "needs_review" and r1.round == 2
    inc3 = next(x for x in r1.criteria or [] if x.id == "GZQO-INC-03")
    assert inc3.logic.value is not None and inc3.logic.value.num == 30  # type: ignore[union-attr]
    r2 = _approve(env, _decisions(env["store"], r1.review_xlsx_key))
    assert r2.status == "approved" and r2.round == 2


def test_max_three_rounds(env: dict[str, Any]) -> None:
    res = _compile(env)
    key = res.review_xlsx_key
    status = ""
    for n in range(3):
        r = _approve(env, _decisions(env["store"], key, {"GZQO-INC-03": ("edit", f"BMI ≥{28 + n} kg/m2")}))
        status = r.status
        key = r.review_xlsx_key or key
    assert status == "max_rounds"


def test_equivalence_failure_blocks_approval(env: dict[str, Any]) -> None:
    env["gate"].failing = ["GZQO-INC-08"]
    res = _compile(env)
    assert "GZQO-INC-08" in (res.needs_review or [])
    out = _approve(env, _decisions(env["store"], res.review_xlsx_key))
    assert out.status == "needs_review" and "GZQO-INC-08" in (out.blocking or [])


def test_reject_removes_criterion_from_compiled_artifacts(env: dict[str, Any]) -> None:
    res = _compile(env)
    out = _approve(env, _decisions(env["store"], res.review_xlsx_key, {"GZQO-EXC-13": ("reject", None)}))
    assert out.status == "approved"
    sql = (env["tmp"] / "repo" / "GZQO" / "sql" / "GZQO.sql").read_text()
    assert "GZQO-EXC-13" not in sql and "GZQO-EXC-10" in sql


def test_unknown_draft_and_criteria(env: dict[str, Any]) -> None:
    with pytest.raises(CompileFailed):
        _approve(env, [])
    _compile(env)
    with pytest.raises(CompileFailed, match="unknown"):
        _approve(env, [c.ReviewDecision(id="GZQO-INC-99", status="approve")])


def test_phi_in_source_forces_local(env: dict[str, Any]) -> None:
    doc = parse(PDF.read_bytes(), PDF.name)
    inc = list(doc.ie_block.inclusion)
    inc[0] = inc[0] + " Example: 王小明先生，男，65歲，病歷號 12345678."
    doc = doc.model_copy(update={"ie_block": doc.ie_block.model_copy(update={"inclusion": inc})})
    env["store"].put("attachments/J/phi.json", c.dump_json(doc).encode())
    res = _compile(env, key="attachments/J/phi.json")
    assert res.phi_guard_hit


def test_incremental_compile_and_diff(env: dict[str, Any]) -> None:
    res = _compile(env)
    _approve(env, _decisions(env["store"], res.review_xlsx_key))
    doc = parse(PDF.read_bytes(), PDF.name)
    exc = list(doc.ie_block.exclusion)
    exc[5] = "HbA1c >9.5% at screening."
    exc.append("Body mass index (BMI) >50 kg/m² at screening.")
    doc = doc.model_copy(update={"ie_block": doc.ie_block.model_copy(update={"exclusion": exc})})
    env["store"].put("attachments/J/v2.json", c.dump_json(doc).encode())
    d = env["comp"].diff(c.DiffRequest(ruleset="GZQO", new_doc_key="attachments/J/v2.json"))
    assert d.changed == ["GZQO-EXC-06"] and len(d.added) == 1 and not d.removed and len(d.unchanged) == 22
    calls = env["gate"].calls
    res2 = _compile(env, key="attachments/J/v2.json")
    assert res2.version == "1.1.0" and len(res2.criteria) == 24
    by_id = {x.id: x for x in res2.criteria}
    assert by_id["GZQO-EXC-06"].text == "HbA1c >9.5% at screening."  # changed criterion keeps its id
    assert by_id["GZQO-EXC-15"].text.startswith("Body mass index")
    assert env["gate"].calls == calls + 1
