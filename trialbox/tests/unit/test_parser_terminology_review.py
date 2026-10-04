from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from criteria_compiler.repo import RulesetRepo
from criteria_compiler.review import logic_summary, parse_review_xlsx, review_xlsx
from criteria_compiler.terminology.mapper import Terminology, ValueSets, valueset
from doc_parser.parser import items, parse
from openpyxl import load_workbook
from tb_contracts import CriterionIR, dump_json

ROOT = Path(__file__).resolve().parents[2]
GOLD = json.loads((ROOT / "services/llm_stub/cassettes/ir_extract/GZQO.json").read_text(encoding="utf-8"))["criteria"]


@pytest.mark.parametrize("name", ["GZQO_protocol_v3.pdf", "GZQO_protocol_v3.docx", "GZQO_questions.xlsx"])
def test_parser_extracts_exact_criteria(name: str) -> None:
    data = (ROOT / "tests/fixtures/protocols" / name).read_bytes()
    doc = parse(data, name)
    assert doc.ie_block.inclusion == [c["text"] for c in GOLD if c["kind"] == "inclusion"]
    assert doc.ie_block.exclusion == [c["text"] for c in GOLD if c["kind"] == "exclusion"]
    assert doc.language == "en" and doc.ie_source == "heuristic"
    assert dump_json(doc) == dump_json(parse(data, name))  # deterministic


def test_parser_zh_rule_text_and_llm_fallback() -> None:
    text = "給付規定\n一、給付條件\n1. 年滿十八歲之類風濕性關節炎病患。\n2. 經兩種以上DMARDs治療\n六個月以上無效。\n二、排除條件\n1. 活動性結核病。\n三、續用條件\n1. DAS28下降1.2以上。\n"
    doc = parse(text.encode(), "rule.txt")
    assert doc.language == "zh-TW"
    assert doc.ie_block.inclusion == ["年滿十八歲之類風濕性關節炎病患。", "經兩種以上DMARDs治療六個月以上無效。"]
    assert doc.ie_block.exclusion == ["活動性結核病。"] and doc.ie_block.renewal == ["DAS28下降1.2以上。"]
    plain = parse(
        b"Some document without criteria headings.",
        "x.txt",
        ie_locator=lambda secs: {"inclusion": ["a"], "exclusion": []},
    )
    assert plain.ie_source == "llm" and plain.ie_block.inclusion == ["a"]


def test_items_split() -> None:
    assert items("1. First line\ncontinued\n2. Second") == ["First line continued", "Second"]
    assert items("(1) 甲\n（2） 乙\n• 丙") == ["甲", "乙", "丙"]


def test_terminology_steps() -> None:
    t = Terminology()
    assert t.map("condition", "Gout").method == "synonym"
    assert t.map("medication", "allopurinol").method == "exact"
    unmapped = t.map("condition", "psoriatic arthritis")
    assert unmapped.method == "none" and unmapped.needs_review
    via_llm = t.map(
        "condition", "type 2 diabetes", llm=lambda n, d, cands: ([(cands[0].code, cands[0].system, 0.7)], True)
    )
    assert via_llm.method == "llm" and via_llm.needs_review
    vs = valueset("X", "VS_SBP", "SBP", t.map("observation", "systolic blood pressure").concepts, "1.0.0")
    assert vs["compose"]["include"][0]["concept"][0]["code"] == "8480-6"
    assert ValueSets("X", {"VS_SBP": vs}, t).is_component_set("VS_SBP")


def test_review_xlsx_roundtrip_and_summary() -> None:
    crit = [
        CriterionIR.model_validate(
            {
                "id": "X-INC-01",
                "ruleset": "X",
                "text": "BMI >= 27",
                "kind": "inclusion",
                "class": "structured",
                "logic": {
                    "domain": "observation",
                    "valueset": "VS_BMI_COMPONENTS",
                    "derived": "bmi",
                    "quantifier": "latest",
                    "value": {"op": ">=", "num": 27, "unit": "kg/m2"},
                },
            }
        )
    ]
    assert logic_summary(crit[0].logic, {}).startswith("Latest BMI (from latest height and weight) within 365 days")
    data = review_xlsx("X", "1.0.0", 1, "J", crit, {}, None)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["review"]
    ws["I2"] = "edit"
    ws["J2"] = "BMI >= 30"
    buf = io.BytesIO()
    wb.save(buf)
    decisions, meta = parse_review_xlsx(buf.getvalue())
    assert decisions[0].status == "edit" and decisions[0].edited_text == "BMI >= 30" and meta["ruleset"] == "X"
    ws["I2"] = "maybe"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        parse_review_xlsx(buf.getvalue())


def test_repo_seed_tags_and_promote(tmp_path: Path) -> None:
    repo = RulesetRepo(tmp_path / "repo", ROOT / "rulesets")
    assert "GZQO/v1.0.0" in repo.tags()  # vendor-shipped approved ruleset is tagged on seeding
    assert (tmp_path / "repo" / "GZQO" / "manifest.yaml").exists()
    repo.write_dir("draft/J", "main", "NEW/", {"manifest.yaml": b"id: NEW\n"}, "draft")
    assert "NEW/manifest.yaml" not in repo.files("main")
    repo.tag("NEW/v1.0.0", "draft/J", "ok")
    repo.promote("NEW", "NEW/v1.0.0", "release")
    assert (tmp_path / "repo" / "NEW" / "manifest.yaml").read_text() == "id: NEW\n"
    with pytest.raises(ValueError):
        repo.tag("NEW/v1.0.0", "main", "dup")
