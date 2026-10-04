"""Phase 3 DoD on the compose test stack: a real email ``FEAS GZQO`` -> GreenMail (IMAP intake) -> mail-gateway ->
orchestrator -> FEAS -> mail-gateway -> MailHog (SMTP sink) with the feasibility PDF; and the APPROVE loop over email
(review.xlsx round trip) for a protocol attached to ``FEAS``.

    make up-test && make test-e2e
"""

from __future__ import annotations

import io
import os
import smtplib
import time
from collections.abc import Callable
from email import message_from_bytes, policy
from email.message import EmailMessage
from email.utils import make_msgid
from pathlib import Path

import httpx
import pytest
from pypdf import PdfReader

pytestmark = pytest.mark.e2e
ROOT = Path(__file__).resolve().parents[2]
MAILHOG = "http://127.0.0.1:8025"
ORCH = "http://127.0.0.1:8010"
AUTH = "mx.hospa.test; spf=pass smtp.mailfrom=hospa.test; dkim=pass header.d=hospa.test; dmarc=pass"


def send(
    subject: str,
    attachments: tuple[tuple[str, bytes], ...] = (),
    sender: str = "crc1@hospa.test",
    in_reply_to: str | None = None,
) -> str:
    m = EmailMessage()
    m["From"] = sender
    m["To"] = "trialbox@hospa.test"
    m["Subject"] = subject
    m["Message-ID"] = mid = make_msgid(domain="hospa.test")
    m["Authentication-Results"] = AUTH  # what the hospital MTA adds in front of the intake mailbox
    if in_reply_to:
        m["In-Reply-To"] = in_reply_to
    m.set_content("TrialBox E2E")
    for name, data in attachments:
        m.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    with smtplib.SMTP("127.0.0.1", 3025, timeout=30) as s:
        s.send_message(m)
    return mid


def outbox() -> list[EmailMessage]:
    items = httpx.get(f"{MAILHOG}/api/v2/messages", params={"limit": 500}, timeout=30).json()["items"]
    out = []
    for it in items:
        msg = message_from_bytes(it["Raw"]["Data"].encode("utf-8", "surrogateescape"), policy=policy.default)
        assert isinstance(msg, EmailMessage)
        out.append(msg)
    return out


def wait_for(pred: Callable[[EmailMessage], bool], timeout: float = 1800) -> EmailMessage:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        for m in outbox():
            if pred(m):
                return m
        time.sleep(3)
    raise AssertionError("expected mail did not arrive")


def attachments(m: EmailMessage) -> dict[str, bytes]:
    out = {}
    for part in m.iter_attachments():
        data = part.get_payload(decode=True)
        assert isinstance(data, bytes)
        out[str(part.get_filename())] = data
    return out


def job_id_from(m: EmailMessage) -> str:
    return str(m["Subject"]).split()[1]


@pytest.fixture(scope="module", autouse=True)
def stack() -> None:
    try:
        assert httpx.get(f"{ORCH}/healthz", timeout=5).status_code == 200
        httpx.get(f"{MAILHOG}/api/v2/messages", timeout=5).raise_for_status()
    except (httpx.HTTPError, AssertionError):
        pytest.skip("compose test stack not running (make up-test)")


def test_feas_email_to_pdf() -> None:
    t0 = time.monotonic()
    tag = f"e2e-{os.getpid()}"
    send(f"FEAS GZQO -- {tag}")
    job: dict[str, object] = {}
    deadline = time.monotonic() + 300
    while not job and time.monotonic() < deadline:  # the gateway polls IMAP every 5 s in the test profile
        jobs = httpx.get(f"{ORCH}/jobs", params={"type": "FEAS", "limit": 50}, timeout=30).json()
        job = next((j for j in jobs if j.get("comment") == tag), {})
        time.sleep(2)
    assert job, "job not created from the email"
    wait_for(lambda m: str(m["Subject"]).startswith(f"Received {job['job_id']}"), 300)
    done = wait_for(lambda m: str(m["Subject"]).startswith(f"Done {job['job_id']}"), 1800)
    elapsed = time.monotonic() - t0
    assert elapsed < 30 * 60  # SPEC phase 3 DoD: PDF within 30 min
    atts = attachments(done)
    pdf = next(v for k, v in atts.items() if k.endswith(".pdf"))
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Eligibility funnel" in text and f"job {job['job_id']}" in text
    assert any(k.endswith(".xlsx") for k in atts) and any(k.endswith(".json") for k in atts)
    final = httpx.get(f"{ORCH}/jobs/{job['job_id']}", timeout=30).json()
    assert final["state"] == "done" and len(final["outputs"]) == 4


def test_approve_loop_by_email() -> None:
    ruleset = f"GZQO-E{os.getpid() % 100000}"
    pdf = (ROOT / "tests/fixtures/protocols/GZQO_protocol_v3.pdf").read_bytes()
    send(f"FEAS {ruleset}", (("GZQO_protocol_v3.pdf", pdf),))
    review = wait_for(lambda m: str(m["Subject"]).startswith(f"APPROVE {ruleset} version=1.0.0"), 1800)
    xlsx = next(v for k, v in attachments(review).items() if k.endswith(".xlsx"))
    send("Approved, thanks", (("review.xlsx", xlsx),), in_reply_to=str(review["Message-ID"]))
    wait_for(lambda m: str(m["Subject"]) == f"Ruleset {ruleset} v1.0.0 approved", 1800)
    feas_id = str(review["Subject"]).rsplit("job ", 1)[1].strip()
    done = wait_for(lambda m: str(m["Subject"]).startswith(f"Done {feas_id}"), 1800)
    assert any(k.endswith(".pdf") for k in attachments(done))


def _unzip(m: EmailMessage) -> dict[str, bytes]:
    import tempfile

    import py7zr

    ((name, data),) = attachments(m).items()
    assert name.endswith(".7z")
    pw = wait_for(lambda x: str(x["Subject"]).startswith("Password") and str(m["Subject"]) in str(x["Subject"]), 120)
    body = pw.get_body(("plain",))
    assert body is not None
    password = next(ln.strip() for ln in str(body.get_content()).splitlines() if ln.startswith("    "))
    with tempfile.TemporaryDirectory() as td, py7zr.SevenZipFile(io.BytesIO(data), password=password) as z:
        z.extractall(td)
        return {p.name: p.read_bytes() for p in Path(td).rglob("*") if p.is_file()}


def test_screen_email_to_encrypted_list() -> None:
    import json

    tag = f"screen-{os.getpid()}"
    send(f"SCREEN GZQO version=1.0.0 -- {tag}")
    job: dict[str, object] = {}
    deadline = time.monotonic() + 300
    while not job and time.monotonic() < deadline:
        jobs = httpx.get(f"{ORCH}/jobs", params={"type": "SCREEN", "limit": 50}, timeout=30).json()
        job = next((j for j in jobs if j.get("comment") == tag), {})
        time.sleep(2)
    assert job
    done = wait_for(lambda m: str(m["Subject"]).startswith(f"Done {job['job_id']}"), 1800)
    assert any(k.endswith(".pdf") for k in attachments(done))
    lst = wait_for(lambda m: str(m["Subject"]) == f"Candidates GZQO v1.0.0 — job {job['job_id']}", 300)
    assert "pi@hospa.test" in str(lst["To"]) and "sponsor" not in str(lst["To"])
    files = _unzip(lst)
    cl = json.loads(next(v for k, v in files.items() if k.endswith(".json")))
    assert cl["rows"] and cl["summary"]["high"] + cl["summary"]["review"] == len(cl["rows"])
    final = httpx.get(f"{ORCH}/jobs/{job['job_id']}", timeout=30).json()
    assert final["state"] == "done" and final["metrics"]["patients_scoped"] > 50


def test_nav_email_to_lists_and_drafts() -> None:
    import docx

    tag = f"nav-{os.getpid()}"
    send(f"NAV RA-BIO dept=RHEU -- {tag}", sender="nurse-rheu@hospa.test")
    job: dict[str, object] = {}
    deadline = time.monotonic() + 300
    while not job and time.monotonic() < deadline:
        jobs = httpx.get(f"{ORCH}/jobs", params={"type": "NAV", "limit": 50}, timeout=30).json()
        job = next((j for j in jobs if j.get("comment") == tag), {})
        time.sleep(2)
    assert job
    lst = wait_for(lambda m: str(m["Subject"]) == f"NAV RA-BIO RHEU — job {job['job_id']}", 1800)
    files = _unzip(lst)
    assert any(k.startswith("nav_lists_RA-BIO_RHEU") and k.endswith(".xlsx") for k in files)
    drafts = [v for k, v in files.items() if k.endswith(".docx")]
    assert drafts
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(drafts[0])).paragraphs)
    assert "附表十五" in text and "[待補]" in text
    final = httpx.get(f"{ORCH}/jobs/{job['job_id']}", timeout=30).json()
    assert final["state"] == "done"


def _job_by_tag(jtype: str, tag: str, timeout: float = 300) -> dict[str, object]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        jobs = httpx.get(f"{ORCH}/jobs", params={"type": jtype, "limit": 50}, timeout=30).json()
        job: dict[str, object] = next((j for j in jobs if j.get("comment") == tag), {})
        if job:
            return job
        time.sleep(2)
    raise AssertionError(f"no {jtype} job tagged {tag}")


def test_nav_onc_twpas_then_submit_dry_run() -> None:
    """Phase 6 on compose: ``NAV ONC-OSI`` -> TWPAS bundles validated by the HL7 validator (0 errors) and pre-checked
    on HAPI; the physician's ``SUBMIT`` reply -> dry-run ClaimResponse, nothing leaves the box."""
    import json

    tag = f"onc-{os.getpid()}"
    send(f"NAV ONC-OSI dept=ONC -- {tag}", sender="nurse-onc@hospa.test")
    job = _job_by_tag("NAV", tag)
    lst = wait_for(lambda m: str(m["Subject"]) == f"NAV ONC-OSI ONC — job {job['job_id']}", 1800)
    files = _unzip(lst)
    rep = json.loads(files["twpas_validation_ONC-OSI.json"])
    assert rep["validator"] == "hl7-validator" and rep["ig"].startswith("tw.gov.mohw.nhi.pas#")
    built = [b for b in rep["bundles"] if b.get("file")]
    assert len(built) >= 10 and all(b["file"] in files for b in built)
    assert all(not b["validator_errors"] for b in built), [b["validator_errors"] for b in built][:3]
    assert all(b["precheck"]["passed"] is True for b in built), [b["precheck"] for b in built][:3]
    pid = str(built[0]["pid"])
    tag2 = f"submit-{os.getpid()}"
    send(f"SUBMIT ONC-OSI pid={pid} bundle={job['job_id']} -- {tag2}", sender="onc-dr@hospa.test")
    sub = _job_by_tag("SUBMIT", tag2)
    done = wait_for(lambda m: str(m["Subject"]).startswith(f"Done {sub['job_id']}") and bool(attachments(m)), 600)
    assert "onc-dr@hospa.test" in str(done["To"])
    receipt = json.loads(_unzip(done)[f"twpas_claimresponse_ONC-OSI_{pid[:12]}.json"])
    assert receipt["resourceType"] == "ClaimResponse" and receipt["outcome"] == "queued"
    assert "DRY RUN" in receipt["disposition"]
    final = httpx.get(f"{ORCH}/jobs/{sub['job_id']}", timeout=30).json()
    assert final["state"] == "done"


def test_cohort_and_alliance_merge(tmp_path: Path) -> None:
    """Phase 7 on compose: ``COHORT GOUT-COH`` on the root box (site A) -> alliance CSV, report PDF and the trial
    simulation for the three recruiting gout trials; site B's table (an in-process member box on the site-B export)
    arrives by email as ``COHORT MERGE`` from trialbox@hospb.test -> merged table with both sites."""
    import csv
    from datetime import date

    from openpyxl import load_workbook

    from tests.inproc_box import SITE_B, ingest_site, make_box

    tag = f"cohort-{os.getpid()}"
    send(f"COHORT GOUT-COH -- {tag}")
    job = _job_by_tag("COHORT", tag)
    done = wait_for(lambda m: str(m["Subject"]).startswith(f"Done {job['job_id']}") and bool(attachments(m)), 1800)
    files = attachments(done)
    table = next(v for k, v in files.items() if k.startswith("cohort_table_GOUT-COH_DEMO-A_"))
    rows = list(csv.DictReader(io.StringIO(table.decode("utf-8"))))
    assert {r["quarter"] for r in rows} == {"2025Q4", "2026Q1", "2026Q2", "2026Q3"}
    wb = load_workbook(io.BytesIO(next(v for k, v in files.items() if k.startswith("trial_sim_"))))
    ncts = [r[0] for r in list(wb["trials"].iter_rows(values_only=True))[1:]]
    assert ncts == ["NCT99000001", "NCT99000002", "NCT99000003"]
    pdf = next(v for k, v in files.items() if k.endswith(".pdf"))
    assert "Cohort report" in "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    # member box B (in process) produces and mails its table to the root
    sec = tmp_path / "sec"
    sec.mkdir()
    (sec / "site_hmac.key").write_bytes(b"trialbox-test-site-key-B-0123456789abcd")
    ingest_site(ROOT / "tests/fixtures/synthetic_patients/site-b", tmp_path / "lake", sec, "2026-10-04", "DEMO-B")
    box_b = make_box(tmp_path / "box", tmp_path / "lake", sec, date.today(), settings_update=SITE_B)
    box_b.mail(
        "COHORT GOUT-COH lookback=1",
        sender="crc1@hospb.test",
        auth="mx.hospb.test; spf=pass smtp.mailfrom=hospb.test; dkim=pass header.d=hospb.test; dmarc=pass",
    )
    box_b.orch.drain()
    share = next(s for s in box_b.sent if s.subject.startswith("COHORT MERGE"))
    ((name, data),) = share.attachments().items()
    tag2 = f"merge-{os.getpid()}"
    m = EmailMessage()
    m["From"] = "trialbox@hospb.test"
    m["To"] = "trialbox@hospa.test"
    m["Subject"] = f"COHORT MERGE -- {tag2}"
    m["Message-ID"] = make_msgid(domain="hospb.test")
    m["Authentication-Results"] = "mx.hospa.test; spf=pass smtp.mailfrom=hospb.test; dkim=pass header.d=hospb.test"
    m.set_content(share.text())
    m.add_attachment(data, maintype="text", subtype="csv", filename=name)
    with smtplib.SMTP("127.0.0.1", 3025, timeout=30) as s:
        s.send_message(m)
    merge = _job_by_tag("COHORT", tag2)
    out = wait_for(lambda x: str(x["Subject"]).startswith(f"Done {merge['job_id']}") and bool(attachments(x)), 600)
    assert "trialbox@hospb.test" in str(out["To"])
    merged = list(csv.DictReader(io.StringIO(next(iter(attachments(out).values())).decode("utf-8"))))
    pop = {r["site_id"]: r for r in merged if r["criterion_id"] == "GOUT-COH-INC-01" and r["quarter"] == "2026Q3"}
    assert set(pop) == {"DEMO-A", "DEMO-B", "ALLIANCE"}
    assert int(pop["ALLIANCE"]["n"]) == int(pop["DEMO-A"]["n"]) + int(pop["DEMO-B"]["n"])
