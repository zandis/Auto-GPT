"""FEEDBACK (SPEC §8.2): the CRC replies ``FEEDBACK <ruleset>`` with ``screen_feedback.csv`` (columns
``pid, outcome, reason_code, note``) or the candidate workbook's ``feedback`` sheet; rows are validated against the
candidate pool and stored in the ``feedback`` table used by ``calibration``."""

from __future__ import annotations

import csv
import io
from typing import Any, cast

from tb_contracts import ScreenFeedbackRow

from orchestrator.clients import StepFailed
from orchestrator.core import Ctx, Outcome

OUTCOMES = ("enrolled", "screen_fail", "declined", "not_contacted")


def read_rows(name: str, data: bytes) -> list[dict[str, Any]]:
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        if "feedback" not in wb.sheetnames:
            raise StepFailed("feedback", f"{name} has no 'feedback' sheet")
        it = wb["feedback"].iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it)]
        return [dict(zip(header, ["" if v is None else str(v) for v in r], strict=False)) for r in it]
    text = data.decode("utf-8-sig")
    return [{(k or "").strip().lower(): (v or "") for k, v in r.items()} for r in csv.DictReader(io.StringIO(text))]


def run(ctx: Ctx) -> Outcome:
    job = ctx.job
    ruleset = (job.ruleset or "").upper()
    files = [f for f in job.inputs or [] if f.filename.lower().endswith((".csv", ".xlsx"))]
    if not files:
        raise StepFailed("feedback", "Attach screen_feedback.csv (pid, outcome, reason_code, note).")
    ctx.state("running")
    pool = {p["pid"] for p in ctx.orch.db.pool(ruleset)}
    good: list[dict[str, Any]] = []
    errors: list[str] = []
    for f in files:
        for i, r in enumerate(read_rows(f.filename, ctx.input_bytes(f)), start=2):
            pid = (r.get("pid") or "").strip()
            outcome = (r.get("outcome") or "").strip().lower()
            if not pid and not outcome:
                continue
            if not outcome:
                continue  # template row not filled in
            if outcome not in OUTCOMES:
                errors.append(f"{f.filename} row {i}: outcome must be one of {', '.join(OUTCOMES)}")
                continue
            if pid not in pool:
                errors.append(f"{f.filename} row {i}: pid not in the {ruleset} candidate pool")
                continue
            row = ScreenFeedbackRow(
                pid=pid,
                outcome=cast(Any, outcome),
                reason_code=(r.get("reason_code") or "").strip() or None,
                note=(r.get("note") or "").strip() or None,
            )
            good.append(row.model_dump())
    if errors and not good:
        raise StepFailed("feedback", "No valid rows: " + "; ".join(errors[:10]))
    ctx.orch.db.feedback_upsert(ruleset, good, job.job_id, ctx.orch._now().isoformat())
    counts = {o: sum(1 for g in good if g["outcome"] == o) for o in OUTCOMES}
    text = (
        f"Feedback for **{ruleset}** stored: {len(good)} rows ("
        + ", ".join(f"{k} {v}" for k, v in counts.items())
        + ")."
        + ("\n\nSkipped:\n" + "\n".join(f"- {e}" for e in errors[:20]) if errors else "")
    )
    return Outcome(summary_md=text)
