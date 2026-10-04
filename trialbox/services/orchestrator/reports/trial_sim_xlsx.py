"""Trial simulation workbook (SPEC §8.3, tag ``aggregate``): sheet ``trials`` (one row per recruiting trial) and
sheet ``steps`` (each trial's compiled criteria with the remaining count after every applied step)."""

from __future__ import annotations

import io

from tb_common.deterministic import normalize_ooxml
from tb_contracts import TrialSimResult

from orchestrator.reports.common import ReportMeta


def render(sims: TrialSimResult, meta: ReportMeta) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.properties.creator = f"TrialBox {meta.site_id}"
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DCE6F1")
    ws = wb.active
    assert ws is not None
    ws.title = "trials"
    head = [
        "nct_id",
        "title",
        "phases",
        "countries",
        "sponsor",
        "last_update",
        "target_enrollment",
        "draft_version",
        "criteria_counted",
        "criteria_total",
        "equivalence_pct",
        "eligible_now",
        "new_per_month",
        "enrol_12m_P10",
        "enrol_12m_P50",
        "enrol_12m_P90",
        "cached",
        "note",
    ]
    ws.append(head)
    for r in sims.rows:
        ws.append(
            [
                r.nct_id,
                r.title,
                ", ".join(r.phases or []),
                ", ".join(r.countries or []),
                r.sponsor or "",
                str(r.last_update),
                r.target_enrollment,
                r.ruleset_version or "",
                r.criteria_applied,
                r.criteria_total,
                r.equivalence_pct,
                r.eligible_now,
                r.new_per_month,
                r.enrol_12m_low,
                r.enrol_12m_mid,
                r.enrol_12m_high,
                "yes" if r.cached else "no",
                r.note or "",
            ]
        )
    st = wb.create_sheet("steps")
    st.append(["nct_id", "criterion_id", "label", "counted", "remaining_after_step"])
    for r in sims.rows:
        for s in r.steps or []:
            st.append([r.nct_id, s.criterion_id, s.label, "yes" if s.applied else "no", s.remaining])
    about = wb.create_sheet("about")
    for line in (
        (
            "source",
            f"ClinicalTrials.gov API v2 ({sims.source}), condition '{sims.condition}', Taiwan/Japan, recruiting",
        ),
        ("run_date", str(sims.run_date)),
        ("snapshot", str(sims.snapshot)),
        ("reach_rate", sims.reach_rate),
        ("accept_rate", sims.accept_rate),
        ("method", "automatic compile of the registry eligibility text (not reviewed) → FEAS counts → simulation"),
        ("footer", meta.footer),
    ):
        about.append(list(line))
    for sheet in (ws, st, about):
        for cell in sheet[1]:
            cell.font = bold
            cell.fill = fill
    buf = io.BytesIO()
    wb.save(buf)
    return normalize_ooxml(buf.getvalue())
