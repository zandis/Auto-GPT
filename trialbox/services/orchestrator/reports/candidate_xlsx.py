"""Candidate workbook (SPEC §9.2, tag ``phi``): sheets ``candidates``, ``criteria``, ``feedback`` and — for the weekly
microbatch — ``changes``. MRNs are resolved from the in-box pid map and only ever rendered for internal recipients."""

from __future__ import annotations

import io
from collections.abc import Mapping
from dataclasses import dataclass

from tb_common.deterministic import normalize_ooxml
from tb_common.ruleset import Ruleset
from tb_contracts import CandidateRow

from orchestrator.reports.common import ReportMeta

SHOWN = {"pass": "pass", "fail": "fail", "unknown": "unknown", "pending_human": "pending"}
OUTCOMES = "enrolled,screen_fail,declined,not_contacted"


@dataclass(frozen=True)
class Change:
    pid: str
    criterion_id: str
    before: str
    after: str
    quote: str = ""


def render(
    rows: list[CandidateRow],
    rs: Ruleset,
    meta: ReportMeta,
    mrns: Mapping[str, str],
    practitioners: Mapping[str, str],
    changes: list[Change] | None = None,
    title: str = "candidates",
) -> bytes:
    from criteria_compiler.review import logic_summary
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    crit = [c for c in rs.active() if c.kind in ("inclusion", "exclusion")]
    wb = Workbook()
    wb.properties.creator = f"TrialBox {meta.site_id}"
    wb.properties.title = f"{title} {rs.id} v{rs.version}"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    changed_fill = PatternFill("solid", fgColor="F8CBAD")
    ws = wb.active
    assert ws is not None
    ws.title = "candidates"
    header = ["MRN", "pid", "tier", "next_appointment", "practitioner", "department"]
    header += [f"C_{c.id}" for c in crit] + ["n_unknown", "actions", "last_evaluated"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = head_fill
    first_crit_col = 7
    for r in rows:
        by_id = {v.id: v for v in r.criteria}
        values: list[object] = [
            mrns.get(r.pid, ""),
            r.pid,
            r.tier,
            r.next_appointment.strftime("%Y-%m-%d %H:%M") if r.next_appointment else "",
            practitioners.get(r.practitioner_id or "", r.practitioner_id or ""),
            r.department or "",
        ]
        values += [SHOWN[by_id[c.id].verdict] if c.id in by_id else "" for c in crit]
        values += [
            r.n_unknown or 0,
            "; ".join(r.actions or []),
            r.last_evaluated.strftime("%Y-%m-%d %H:%M") if r.last_evaluated else "",
        ]
        ws.append(values)
        rown = ws.max_row
        for j, c in enumerate(crit):
            v = by_id.get(c.id)
            cell = ws.cell(row=rown, column=first_crit_col + j)
            if v is not None and v.evidence is not None and (v.evidence.quote or v.evidence.reason):
                text = f"{v.evidence.quote_date or ''} {v.evidence.quote or ''}".strip() or ""
                if v.evidence.reason:
                    text += f"\n({v.evidence.reason})"
                if v.evidence.confidence is not None:
                    text += f"\nconfidence {v.evidence.confidence:.2f}"
                cell.comment = Comment(text[:1000], "TrialBox")
            if c.id in set(r.changed or []):
                cell.fill = changed_fill
    last = get_column_letter(ws.max_column)
    if ws.max_row > 1:
        rng = f"A2:{last}{ws.max_row}"
        for tier_name, color in (("high", "C6EFCE"), ("review", "FFEB9C")):
            rule = FormulaRule(formula=[f'$C2="{tier_name}"'], fill=PatternFill("solid", fgColor=color))  # type: ignore[no-untyped-call]
            ws.conditional_formatting.add(rng, rule)
    ws.freeze_panes = "C2"
    for i, w in enumerate([12, 34, 8, 18, 22, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([])
    ws.append([meta.footer])

    cs = wb.create_sheet("criteria")
    cs.append(["id", "kind", "class", "label", "protocol_text", "how_computed", "action"])
    for cell in cs[1]:
        cell.font = bold
        cell.fill = head_fill
    for c in crit:
        how = (
            logic_summary(c.logic, rs.valuesets)
            if c.class_ == "structured"
            else (c.note_question or "" if c.class_ == "note" else "ask at screening")
        )
        cs.append([c.id, c.kind, c.class_, c.label or "", c.text, how, c.action or ""])
    for col, w in zip("ABCDEFG", (16, 10, 10, 22, 60, 60, 30), strict=True):
        cs.column_dimensions[col].width = w

    fb = wb.create_sheet("feedback")
    fb.append(["pid", "MRN", "outcome", "reason_code", "note"])
    for cell in fb[1]:
        cell.font = bold
        cell.fill = head_fill
    for r in rows:
        fb.append([r.pid, mrns.get(r.pid, ""), "", "", ""])
    dv = DataValidation(type="list", formula1=f'"{OUTCOMES}"', allow_blank=True)
    fb.add_data_validation(dv)
    if rows:
        dv.add(f"C2:C{len(rows) + 1}")
    for col, w in zip("ABCDE", (34, 12, 16, 16, 40), strict=True):
        fb.column_dimensions[col].width = w

    if changes is not None:
        ch = wb.create_sheet("changes")
        ch.append(["MRN", "pid", "criterion", "before", "after", "evidence"])
        for cell in ch[1]:
            cell.font = bold
            cell.fill = head_fill
        for x in changes:
            ch.append(
                [
                    mrns.get(x.pid, ""),
                    x.pid,
                    x.criterion_id,
                    SHOWN.get(x.before, x.before),
                    SHOWN.get(x.after, x.after),
                    x.quote,
                ]
            )
        for col, w in zip("ABCDEF", (12, 34, 16, 10, 10, 60), strict=True):
            ch.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return normalize_ooxml(buf.getvalue())
