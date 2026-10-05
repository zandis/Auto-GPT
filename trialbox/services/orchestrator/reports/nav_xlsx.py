"""NAV lists workbook (SPEC §9.3, tag ``phi``): one workbook per department with sheets ``likely_eligible``,
``renewal_due``, ``doc_gaps``, ``maybe_ineligible``; rows in appointment order."""

from __future__ import annotations

import io
from collections.abc import Mapping

from tb_common.deterministic import normalize_ooxml
from tb_common.ruleset import Ruleset
from tb_contracts import NavRow

from orchestrator.reports.common import ReportMeta

SHEETS = ("likely_eligible", "renewal_due", "doc_gaps", "maybe_ineligible")
SHOWN = {"pass": "pass", "fail": "fail", "unknown": "unknown", "pending_human": "pending"}


def render(
    lists: Mapping[str, list[NavRow]],
    rs: Ruleset,
    meta: ReportMeta,
    mrns: Mapping[str, str],
    practitioners: Mapping[str, str],
    files: Mapping[str, str],
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    crit = [c for c in rs.active()]
    wb = Workbook()
    wb.properties.creator = f"TrialBox {meta.site_id}"
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DCE6F1")
    first = True
    for name in SHEETS:
        ws = wb.active if first else wb.create_sheet(name)
        assert ws is not None
        ws.title = name
        first = False
        header = ["MRN", "pid", "next_appointment", "practitioner", "approval_end"]
        header += [f"C_{c.id}" for c in crit]
        header += ["missing_items", "suggested_orders", "draft_file", "bundle_file", "precheck_result"]
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
            cell.fill = fill
        for r in lists.get(name, []):
            by_id = {v.id: v for v in r.criteria}
            pre = r.precheck
            pre_txt = (
                ""
                if pre is None or pre.passed is None
                else ("passed" if pre.passed else "failed: " + "; ".join(pre.issues or []))
            )
            ws.append(
                [
                    mrns.get(r.pid, ""),
                    r.pid,
                    r.next_appointment.strftime("%Y-%m-%d %H:%M") if r.next_appointment else "",
                    practitioners.get(r.practitioner_id or "", r.practitioner_id or ""),
                    r.approval_end.isoformat() if r.approval_end else "",
                ]
                + [SHOWN[by_id[c.id].verdict] if c.id in by_id else "" for c in crit]
                + [
                    "; ".join(m.item for m in r.missing or []),
                    "; ".join(m.suggested_order for m in r.missing or [] if m.suggested_order),
                    files.get(f"draft:{r.pid}", ""),
                    files.get(f"bundle:{r.pid}", ""),
                    pre_txt,
                ]
            )
            row = ws.max_row
            for j, c in enumerate(crit):
                v = by_id.get(c.id)
                if v is not None and v.evidence is not None and v.evidence.quote:
                    ws.cell(row=row, column=6 + j).comment = Comment(
                        f"{v.evidence.quote_date or ''} {v.evidence.quote}"[:1000], "TrialBox"
                    )
        ws.freeze_panes = "C2"
        ws.append([])
        ws.append([meta.footer])
    buf = io.BytesIO()
    wb.save(buf)
    return normalize_ooxml(buf.getvalue())
