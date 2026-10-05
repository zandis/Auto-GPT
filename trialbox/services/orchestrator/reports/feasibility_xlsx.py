"""Feasibility Excel companion (SPEC §9.1): sheets ``funnel``, ``sensitivity``, ``monthly``, ``criteria``,
``assumptions``. Byte-reproducible (fixed zip timestamps and core properties)."""

from __future__ import annotations

import io

from tb_common.deterministic import normalize_ooxml
from tb_common.ruleset import Ruleset
from tb_contracts import FeasibilityResult

from orchestrator.reports.common import ReportMeta


def render(result: FeasibilityResult, rs: Ruleset, meta: ReportMeta) -> bytes:
    from criteria_compiler.review import codes_text, logic_summary
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.properties.creator = f"TrialBox {meta.site_id}"
    wb.properties.title = f"Feasibility {rs.id} v{rs.version}"
    wb.properties.description = meta.footer
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")

    def sheet(title: str, header: list[str], rows: list[list[object]], widths: list[int], first: bool = False) -> None:
        ws = wb.active if first else wb.create_sheet(title)
        assert ws is not None
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = w
        ws.freeze_panes = "A2"
        ws.append([])
        ws.append([meta.footer])

    funnel_rows: list[list[object]] = [["START", result.population, result.start_n, None, 100.0, None, True]]
    for f in result.funnel:
        funnel_rows.append([f.criterion_id, f.label, f.remaining, f.dropped, f.pct, f.unknown, f.applied is not False])
    sheet(
        "funnel",
        ["criterion_id", "label", "remaining", "dropped", "pct_of_start", "unknown", "applied"],
        funnel_rows,
        [16, 40, 12, 12, 12, 12, 10],
        first=True,
    )
    sheet(
        "sensitivity",
        ["criterion_id", "variant", "remaining", "delta_vs_base"],
        [[s.criterion_id, s.variant, s.remaining, s.delta] for s in result.sensitivity],
        [16, 22, 12, 14],
    )
    sheet("monthly", ["month", "new_eligible"], [[m.month, m.n] for m in result.monthly_new], [10, 14])
    crit_rows: list[list[object]] = []
    for c in rs.active():
        structured = c.class_ == "structured"
        crit_rows.append(
            [
                c.id,
                c.kind,
                c.class_,
                c.label or "",
                c.text,
                logic_summary(c.logic, rs.valuesets) if structured else "not applied in counts",
                codes_text(c.logic, rs.valuesets) if structured else "",
                c.fallback or "",
            ]
        )
    sheet(
        "criteria",
        ["id", "kind", "class", "label", "protocol_text", "how_computed", "codes", "fallback"],
        crit_rows,
        [16, 10, 10, 22, 60, 60, 60, 30],
    )
    sim = result.simulation
    assumptions: list[list[object]] = [
        ["reach_rate", sim.reach_rate, sim.source],
        ["accept_rate", sim.accept_rate, sim.source],
        ["capacity_per_month", sim.capacity_per_month, "settings.practitioners"],
        ["monthly_eligible_mean", sim.monthly_eligible_mean, "funnel (last 12 months)"],
        ["months", sim.months, "SPEC §8.1"],
        ["iterations", sim.iterations, "simulation"],
        ["enrolment_low_p10", sim.low, "simulation"],
        ["enrolment_mid_p50", sim.mid, "simulation"],
        ["enrolment_high_p90", sim.high, "simulation"],
        ["snapshot", meta.snapshot, "lake"],
        ["run_date", meta.run_date, "job"],
        ["lookback_months", result.lookback_months, "options/manifest"],
    ]
    assumptions += [[f"note_{i + 1}", n, ""] for i, n in enumerate(result.notes)]
    sheet("assumptions", ["parameter", "value", "source"], assumptions, [26, 60, 26])
    buf = io.BytesIO()
    wb.save(buf)
    return normalize_ooxml(buf.getvalue())
