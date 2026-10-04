"""Review package (SPEC §9.6): ``review.html`` (inline), ``review.xlsx`` (decision columns), draft zip.

``review.xlsx`` sheet ``review`` columns: id, kind, class, text, summary, codes, equivalence_pct, flags, decision,
edited_text, edited_class, comment. ``decision`` is pre-filled with ``approve``; the reviewer changes it to
``reject`` or ``edit``. :func:`parse_review_xlsx` turns a returned workbook into ``ReviewDecision`` models.
"""

from __future__ import annotations

import html
import io
import json
import zipfile
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from tb_contracts import Atom, BoolExpr, CriterionIR, EquivalenceReport, ReviewDecision

from criteria_compiler.semantics import window

COLUMNS = [
    "id",
    "kind",
    "class",
    "text",
    "summary",
    "codes",
    "equivalence_pct",
    "flags",
    "decision",
    "edited_text",
    "edited_class",
    "comment",
]
OP_TEXT = {">": ">", ">=": "≥", "<": "<", "<=": "≤", "=": "=", "between": "between", "in": "is one of"}


def _window_text(atom: Atom) -> str:
    w = window(atom)
    if w.from_days is None:
        return "ever (up to the index date)" if w.to_days == 0 else f"up to {abs(w.to_days)} days before index"
    lo = abs(w.from_days)
    if w.to_days == 0:
        return f"within {lo} days before the index date"
    return f"between {lo} and {abs(w.to_days)} days before the index date"


def _title(atom: Atom, valuesets: dict[str, dict[str, Any]]) -> str:
    vs = valuesets.get(atom.valueset) or {}
    return str(vs.get("title") or atom.valueset)


def atom_summary(atom: Atom, valuesets: dict[str, dict[str, Any]]) -> str:
    if atom.domain == "demographic":
        parts = []
        if atom.age and atom.age.min is not None:
            parts.append(f"age ≥ {atom.age.min:g}")
        if atom.age and atom.age.max is not None:
            parts.append(f"age ≤ {atom.age.max:g}")
        if atom.sex:
            parts.append(f"sex = {atom.sex}")
        return " and ".join(parts) or "patient record exists"
    title = _title(atom, valuesets)
    win = _window_text(atom)
    if atom.derived:
        label = {
            "bmi": "BMI (from latest height and weight)",
            "egfr": "eGFR CKD-EPI 2021 (from latest creatinine)",
            "das28": "DAS28 (TJC28, SJC28, patient global, ESR or CRP)",
            "basdai": "BASDAI",
            "asdas": "ASDAS-CRP",
        }[atom.derived]
        if atom.value and atom.value.op:
            return f"Latest {label} {win} {OP_TEXT[atom.value.op]} {_num(atom)}"
        return f"{label} computable {win}"
    if atom.duration and atom.duration.min_days:
        gap = 30 if atom.duration.gap_days is None else atom.duration.gap_days
        return f"Continuous {title} for ≥ {atom.duration.min_days} days {win} (gaps ≤ {gap} days allowed)"
    if atom.value and atom.value.op:
        q = atom.quantifier or "any"
        lead = {"latest": "Latest", "any": "Any", "all": "Every", "none": "No", "count>=": f"At least {atom.count}"}[q]
        return f"{lead} {title} {win} {OP_TEXT[atom.value.op]} {_num(atom)}"
    q = atom.quantifier or "any"
    noun = {
        "condition": "diagnosis of",
        "medication": "prescription of",
        "procedure": "procedure:",
        "encounter": "visit:",
        "claim": "NHI application for",
        "report": "report:",
        "observation": "result of",
    }
    if q == "count>=":
        return f"At least {atom.count} × {noun[atom.domain]} {title} {win}"
    if q == "none":
        return f"No {noun[atom.domain]} {title} {win}"
    return f"Any {noun[atom.domain]} {title} {win}"


def _num(atom: Atom) -> str:
    v = atom.value
    assert v is not None
    unit = f" {v.unit}" if v.unit else ""
    if v.op == "in":
        return ", ".join(v.codes or [])
    if v.op == "between":
        return f"{v.num:g}–{v.num2:g}{unit}"
    return f"{v.num:g}{unit}"


def logic_summary(expr: Atom | BoolExpr, valuesets: dict[str, dict[str, Any]]) -> str:
    if isinstance(expr, Atom):
        return atom_summary(expr, valuesets)
    parts = [logic_summary(a, valuesets) for a in expr.args]
    if expr.op == "not":
        return f"NOT ({parts[0]})"
    return f" {expr.op.upper()} ".join(f"({p})" if " AND " in p or " OR " in p else p for p in parts)


def codes_text(expr: Atom | BoolExpr, valuesets: dict[str, dict[str, Any]]) -> str:
    from criteria_compiler.semantics import atoms

    seen: list[str] = []
    for a in atoms(expr):
        vs = valuesets.get(a.valueset)
        if not vs:
            continue
        codes = [
            f"{c['code']} {c.get('display', '')}".strip()
            for inc in vs["compose"]["include"]
            for c in inc.get("concept", [])
        ]
        tag = (
            " [needs review]"
            if any(t.get("code") == "needs_review" for t in (vs.get("meta") or {}).get("tag", []))
            else ""
        )
        seen.append(f"{a.valueset}{tag}: " + ("; ".join(codes) if codes else "(no codes)"))
    return "\n".join(dict.fromkeys(seen))


def rows(
    criteria: list[CriterionIR], valuesets: dict[str, dict[str, Any]], tests: EquivalenceReport | None
) -> list[dict[str, Any]]:
    out = []
    for c in criteria:
        eq = (tests.per_criterion.get(c.id) if tests else None) if c.class_ == "structured" else None
        summary = (
            logic_summary(c.logic, valuesets)
            if c.class_ == "structured"
            else (c.note_question or c.human_question or "")
        )
        out.append(
            {
                "id": c.id,
                "kind": c.kind,
                "class": c.class_,
                "text": c.text,
                "summary": summary,
                "codes": codes_text(c.logic, valuesets) if c.class_ == "structured" else "",
                "equivalence_pct": eq,
                "flags": ", ".join(c.flags or []),
                "decision": "reject" if (c.review and c.review.status == "rejected") else "approve",
                "fallback": c.fallback or "",
                "time_sensitive": bool(c.time_sensitive),
            }
        )
    return out


def review_html(
    ruleset: str,
    version: str,
    round_no: int,
    criteria: list[CriterionIR],
    valuesets: dict[str, dict[str, Any]],
    tests: EquivalenceReport | None,
    notes: list[str],
) -> str:
    esc = html.escape
    trs = []
    for r in rows(criteria, valuesets, tests):
        flag = ' style="background:#fff3cd"' if r["flags"] else ""
        eq = "" if r["equivalence_pct"] is None else f"{r['equivalence_pct']:.1f}%"
        trs.append(
            f"<tr{flag}><td>{esc(r['id'])}</td><td>{esc(r['kind'])}</td><td>{esc(r['class'])}</td>"
            f"<td>{esc(r['text'])}</td><td>{esc(r['summary'])}</td>"
            f"<td><pre>{esc(r['codes'])}</pre></td><td>{esc(r['fallback'])}</td><td>{eq}</td>"
            f"<td>{esc(r['flags'])}</td></tr>"
        )
    overall = f"{tests.overall_pct:.1f}% on {tests.sample_size} patients" if tests else "not run"
    failing = ", ".join(tests.failing) if tests and tests.failing else "none"
    notes_html = "".join(f"<li>{esc(n)}</li>" for n in notes)
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{esc(ruleset)} v{esc(version)} review</title>
<style>body{{font-family:'Noto Sans TC',Arial,sans-serif;font-size:13px}}table{{border-collapse:collapse}}
td,th{{border:1px solid #bbb;padding:4px;vertical-align:top}}th{{background:#eee}}pre{{margin:0;white-space:pre-wrap}}</style>
</head><body><h2>Ruleset {esc(ruleset)} v{esc(version)} — review round {round_no}</h2>
<p>CQL ≡ SQL equivalence: {esc(overall)}; failing criteria: {esc(failing)}.</p><ul>{notes_html}</ul>
<p>Reply <b>APPROVE {esc(ruleset)} version={esc(version)}</b> with the attached review.xlsx (decision column).</p>
<table><tr><th>id</th><th>kind</th><th>class</th><th>source text</th><th>how it is computed</th><th>codes</th>
<th>fallback</th><th>CQL≡SQL</th><th>flags</th></tr>{"".join(trs)}</table></body></html>"""


def review_xlsx(
    ruleset: str,
    version: str,
    round_no: int,
    job_id: str,
    criteria: list[CriterionIR],
    valuesets: dict[str, dict[str, Any]],
    tests: EquivalenceReport | None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "review"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in rows(criteria, valuesets, tests):
        ws.append([r.get(c, "") for c in COLUMNS])
    n = ws.max_row
    dv = DataValidation(type="list", formula1='"approve,reject,edit"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"structured,note,human"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    dv.add(f"I2:I{n}")
    dv2.add(f"K2:K{n}")
    widths = {
        "A": 16,
        "B": 11,
        "C": 11,
        "D": 60,
        "E": 60,
        "F": 50,
        "G": 10,
        "H": 22,
        "I": 10,
        "J": 50,
        "K": 12,
        "L": 30,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[7].value:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFF3CD")
    ws.freeze_panes = "B2"
    vs_ws = wb.create_sheet("valuesets")
    vs_ws.append(["valueset", "title", "system", "code", "display"])
    for name, vs in valuesets.items():
        for inc in vs["compose"]["include"]:
            for c in inc.get("concept", []):
                vs_ws.append([name, vs.get("title", ""), inc["system"], c["code"], c.get("display", "")])
    about = wb.create_sheet("about")
    for k, v in [
        ("ruleset", ruleset),
        ("version", version),
        ("round", round_no),
        ("compile_job", job_id),
        (
            "instructions",
            "Set decision to approve / reject / edit. For edit, give edited_text and/or "
            "edited_class. Reply 'APPROVE <ruleset> version=<version>' with this file.",
        ),
    ]:
        about.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_review_xlsx(data: bytes) -> tuple[list[ReviewDecision], dict[str, str]]:
    """Decisions from a returned review workbook, plus its ``about`` metadata."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "review" not in wb.sheetnames:
        raise ValueError("review.xlsx has no 'review' sheet")
    ws = wb["review"]
    header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    missing = [c for c in ("id", "decision") if c not in header]
    if missing:
        raise ValueError(f"review.xlsx is missing columns {missing}")
    idx = {h: i for i, h in enumerate(header)}
    out: list[ReviewDecision] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = row[idx["id"]]
        if not cid:
            continue

        def get(col: str, r: tuple[Any, ...] = row) -> str | None:
            i = idx.get(col)
            v = r[i] if i is not None and i < len(r) else None
            return None if v is None or str(v).strip() == "" else str(v).strip()

        decision = (get("decision") or "").lower()
        if decision not in ("approve", "reject", "edit"):
            raise ValueError(f"{cid}: decision must be approve, reject or edit (got {decision!r})")
        edited_class = get("edited_class")
        if edited_class and edited_class not in ("structured", "note", "human"):
            raise ValueError(f"{cid}: edited_class must be structured, note or human")
        out.append(
            ReviewDecision(
                id=str(cid),
                status=decision,  # type: ignore[arg-type]
                edited_text=get("edited_text"),
                edited_class=edited_class,  # type: ignore[arg-type]
                comment=get("comment"),
            )
        )
    meta: dict[str, str] = {}
    if "about" in wb.sheetnames:
        for k, v in wb["about"].iter_rows(values_only=True):
            if k:
                meta[str(k)] = "" if v is None else str(v)
    return out, meta


def draft_zip(files: dict[str, bytes], extra: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in sorted({**files, **extra}):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))  # reproducible archive
            info.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(info, {**files, **extra}[name])
    return buf.getvalue()


def json_bytes(obj: Any) -> bytes:
    return (json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=True) + "\n").encode("utf-8")
