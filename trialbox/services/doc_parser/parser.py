"""Protocol / rule-document parsing (SPEC §4.2) into ``ParsedDoc``.

Backends produce a list of markdown-like lines (``# heading`` lines + text + ``|`` table rows):

* ``docling`` (production image; DECISIONS D-02): layout + table recognition, exported to markdown;
* ``lite`` (CPU CI / fallback): pypdf text, python-docx paragraphs with heading styles, openpyxl sheets, plain text.

The same deterministic section and I/E heuristics then run on those lines (``Inclusion Criteria|納入條件|選択基準``,
``Exclusion|排除條件|除外基準``, plus ``續用``/renewal and ``應檢附``/documentation for NHI rules). If no block is
found, the ``ie_locate`` LLM prompt is used. Scanned PDFs (< 50 chars/page) go to PaddleOCR when installed.
"""

from __future__ import annotations

import importlib.util
import io
import os
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any, Literal

from tb_common.crypto import sha256_bytes
from tb_contracts import DocSection, DocTable, IEBlock, ParsedDoc

HEADINGS = {
    "inclusion": re.compile(
        r"inclusion criteria|納入條件|納入標準|收案條件|選択基準|適用條件|給付條件|申請條件|世代定義|計數條件", re.I
    ),  # 世代定義 / 計數條件: alliance cohort definitions (population + counted characteristics, D-71)
    "exclusion": re.compile(r"exclusion criteria|排除條件|排除標準|除外基準|不予給付|不得申請", re.I),
    "renewal": re.compile(r"renewal|continuation criteria|續用條件|續用|繼續使用|延長使用|継続基準", re.I),
    "documentation": re.compile(
        r"required documents|documentation requirements|應檢附|檢附資料|申請文件|應附資料", re.I
    ),
}
_ENUM = re.compile(r"^\s*(?:\(?(\d{1,3})[.)、．]|\((\d{1,3})\)|（(\d{1,3})）|[一二三四五六七八九十]+[、.]|[•\-*‧])\s*")
_NUM_HEADING = re.compile(r"^(\d+(?:\.\d+){0,3})\.?\s+(\S.{0,80})$")
_ZH_HEADING = re.compile(r"^(?:第?[一二三四五六七八九十]+[章節、]|[（(][一二三四五六七八九十]+[)）])\s*\S.{0,40}$")
_CJK = re.compile(r"[一-鿿]")
_KANA = re.compile(r"[぀-ヿ]")

IeLocator = Callable[[list[dict[str, Any]]], dict[str, list[str]]]


@dataclass
class _Doc:
    lines: list[str]
    tables: list[DocTable] = field(default_factory=list)
    pages: int = 0
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)
    parser: str = "lite"


def _kind(filename: str, data: bytes) -> str:
    name = filename.lower()
    if data[:4] == b"%PDF" or name.endswith(".pdf"):
        return "pdf"
    if name.endswith(".docx"):
        return "docx"
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    return "text"


# ----------------------------------------------------------------------------- backends
def _lite_pdf(data: bytes) -> _Doc:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    lines: list[str] = []
    chars = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        chars += len(text.strip())
        for raw in text.splitlines():
            line = raw.strip()
            if not line:
                continue
            m = _NUM_HEADING.match(line)
            # "5.1 Inclusion Criteria" / "6 Study Intervention" are headings; "1. text" is an enumerated item
            numbered_item = m is not None and "." not in m.group(1) and line[len(m.group(1)) :].startswith(".")
            if m and not numbered_item and not line.endswith((".", "。", ",", ";")):
                level = m.group(1).count(".") + 1
                lines.append("#" * level + " " + line)
            elif _ZH_HEADING.match(line) or any(p.fullmatch(line.strip(" ：:")) for p in HEADINGS.values()):
                lines.append("## " + line)
            else:
                lines.append(line)
    doc = _Doc(lines, pages=len(reader.pages))
    if reader.pages and chars / len(reader.pages) < 50:
        doc.warnings.append("low text density: scanned PDF suspected")
        if importlib.util.find_spec("paddleocr") is not None:  # pragma: no cover - production image only
            doc.lines, doc.ocr_used = _ocr_pdf(data), True
        else:
            doc.warnings.append("PaddleOCR not installed; OCR skipped")
    return doc


def _ocr_pdf(data: bytes) -> list[str]:  # pragma: no cover - requires paddleocr + pdf rasteriser
    from paddleocr import PaddleOCR
    from pypdf import PdfReader

    ocr = PaddleOCR(lang="chinese_cht", show_log=False)
    out: list[str] = []
    for page in PdfReader(io.BytesIO(data)).pages:
        for image in page.images:
            for line in ocr.ocr(image.data, cls=False)[0] or []:
                out.append(str(line[1][0]))
    return out


def _lite_docx(data: bytes) -> _Doc:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        m = re.match(r"Heading (\d)", style)
        if style == "Title":
            lines.append("# " + text)
        elif m:
            lines.append("#" * int(m.group(1)) + " " + text)
        else:
            lines.append(text)
    tables = [DocTable(caption=None, rows=[[c.text.strip() for c in r.cells] for r in t.rows]) for t in d.tables]
    return _Doc(lines, tables, pages=1)


def _lite_xlsx(data: bytes) -> _Doc:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    tables: list[DocTable] = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        tables.append(DocTable(caption=ws.title, rows=rows))
        header = [h.lower() for h in rows[0]] if rows else []
        sec_i = next((i for i, h in enumerate(header) if h in ("section", "類別", "分類")), None)
        q_i = next((i for i, h in enumerate(header) if "question" in h or "criterion" in h or "條件" in h), None)
        if sec_i is not None and q_i is not None:
            current = None
            for r in rows[1:]:
                if r[sec_i] != current:
                    current = r[sec_i]
                    lines.append(f"## {current}")
                lines.append(f"{len(lines)}. {r[q_i]}")
        else:
            lines.append(f"# {ws.title}")
            lines.extend(" ".join(c for c in r if c) for r in rows)
    return _Doc(lines, tables, pages=len(wb.worksheets))


def _text(data: bytes) -> _Doc:
    text = data.decode("utf-8-sig", errors="replace")
    lines = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        zh_numbered = re.match(r"^[一二三四五六七八九十]+[、.]", line) is not None
        if (
            any(p.search(line) for p in HEADINGS.values()) and len(line) < 60 and (zh_numbered or not _ENUM.match(line))
        ) or (_ZH_HEADING.match(line) and len(line) < 40):
            lines.append("## " + line)
        else:
            lines.append(line)
    return _Doc(lines, pages=1)


def _docling(data: bytes, filename: str) -> _Doc:  # pragma: no cover - production image only
    import tempfile
    from pathlib import Path

    from docling.document_converter import DocumentConverter

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / Path(filename).name
        path.write_bytes(data)
        result = DocumentConverter().convert(str(path))
    md = result.document.export_to_markdown()
    lines = [ln.strip() for ln in md.splitlines() if ln.strip()]
    tables = []
    body = []
    rows: list[list[str]] = []
    for ln in lines:
        if ln.startswith("|"):
            cells = [c.strip() for c in ln.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                rows.append(cells)
            continue
        if rows:
            tables.append(DocTable(caption=None, rows=rows))
            rows = []
        body.append(ln.replace("- ", "", 1) if ln.startswith("- ") else ln)
    if rows:
        tables.append(DocTable(caption=None, rows=rows))
    return _Doc(body, tables, pages=getattr(result.document, "num_pages", lambda: 0)() or 0, parser="docling")


# ----------------------------------------------------------------------------- structure
def sections(lines: list[str]) -> list[DocSection]:
    out: list[DocSection] = []
    title, level = "(preamble)", 0
    buf: list[str] = []
    for ln in lines:
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            if buf or title != "(preamble)":
                out.append(DocSection(title=title, level=level, text="\n".join(buf)))
            title, level, buf = m.group(2).strip(), len(m.group(1)), []
        else:
            buf.append(ln)
    out.append(DocSection(title=title, level=level, text="\n".join(buf)))
    return [s for s in out if s.text or s.title != "(preamble)"]


def items(text: str) -> list[str]:
    """Split a section into criteria: enumerated / bulleted lines start items; other lines continue them."""
    result: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = _ENUM.match(line)
        if m:
            rest = line[m.end() :].strip()
            result.append(rest)
        elif result:
            joiner = "" if _CJK.search(result[-1][-1:] or "") else " "
            result[-1] = (result[-1] + joiner + line).strip()
        else:
            result.append(line)
    return [re.sub(r"\s+", " ", r).strip() for r in result if r.strip()]


def ie_block(secs: list[DocSection]) -> IEBlock:
    found: dict[str, list[str]] = {k: [] for k in HEADINGS}
    refs: dict[str, str] = {}
    for s in secs:
        for key, pat in HEADINGS.items():
            if pat.search(s.title):
                for i, it in enumerate(items(s.text), start=1):
                    found[key].append(it)
                    refs[it] = f"{s.title} #{i}"
                break
    return IEBlock(
        inclusion=found["inclusion"],
        exclusion=found["exclusion"],
        renewal=found["renewal"] or None,
        documentation=found["documentation"] or None,
        refs=refs or None,
    )


def language(lines: list[str]) -> str:
    text = "".join(lines)
    if not text:
        return "und"
    if _KANA.search(text):
        return "ja"
    return "zh-TW" if len(_CJK.findall(text)) / len(text) > 0.1 else "en"


def parse(data: bytes, filename: str, ie_locator: IeLocator | None = None, backend: str | None = None) -> ParsedDoc:
    backend = backend or os.environ.get("TB_PARSER", "auto")
    kind = _kind(filename, data)
    use_docling = (
        kind in ("pdf", "docx") and backend in ("auto", "docling") and importlib.util.find_spec("docling") is not None
    )
    if use_docling:  # pragma: no cover - production image only
        doc = _docling(data, filename)
    elif kind == "pdf":
        doc = _lite_pdf(data)
    elif kind == "docx":
        doc = _lite_docx(data)
    elif kind == "xlsx":
        doc = _lite_xlsx(data)
    else:
        doc = _text(data)
    secs = sections(doc.lines)
    block = ie_block(secs)
    source: Literal["heuristic", "llm", "none"] = "heuristic" if (block.inclusion or block.exclusion) else "none"
    if source == "none" and ie_locator is not None:
        located = ie_locator([{"title": s.title, "text": s.text[:4000]} for s in secs][:60])
        block = IEBlock(inclusion=located.get("inclusion", []), exclusion=located.get("exclusion", []))
        source = "llm" if (block.inclusion or block.exclusion) else "none"
    return ParsedDoc(
        source_sha256=sha256_bytes(data),
        filename=filename,
        parser="docling" if use_docling else "lite",
        language=language(doc.lines),
        pages=doc.pages,
        ocr_used=doc.ocr_used,
        sections=secs,
        tables=doc.tables,
        ie_block=block,
        ie_source=source,
        warnings=doc.warnings or None,
    )
