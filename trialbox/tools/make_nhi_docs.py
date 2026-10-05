"""Synthetic NHI reimbursement-rule sources for the RA-BIO ruleset (fictional; modelled on the public structure of the
NHI drug-benefit rule for RA biologics and its application form 附表十五).

Writes from one definition (RULES):
* ``tests/fixtures/protocols/RA-BIO_給付規定_2026.pdf`` / ``.docx`` — the rule document (zh-TW) the compiler reads;
* ``services/llm_stub/cassettes/ir_extract/RA-BIO.json`` — the gold extraction (stub cassette, D-39);
* ``services/orchestrator/reports/templates/RA-BIO/附表十五.docx`` — the docxtpl application template (SPEC §9.4).

    python tools/make_nhi_docs.py && python tools/make_nhi_docs.py --check   # parse check of the PDF
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "libs"), str(ROOT / "services")]
FONT = ROOT / ".cache" / "fonts"

TITLE = "全民健康保險藥品給付規定 8.2.4 類風濕關節炎生物製劑（合成範例）"
FRONT = [
    ("8.2 免疫製劑", "本文件為 TrialBox 測試用之虛構範例，依公開給付規定之結構撰寫，並非健保署正式公告內容。"),
    (
        "8.2.4 類風濕關節炎生物製劑",
        "適用藥品：adalimumab、etanercept、tocilizumab。申請時應填具附表十五並檢附相關資料。",
    ),
]

DAS = {"domain": "observation", "derived": "das28", "quantifier": "latest"}
# (section, text, label, logic, extras)
RULES: list[tuple[str, str, str, dict[str, Any], dict[str, Any]]] = [
    ("inclusion", "病患年齡須滿十八歲。", "年齡≥18", {"domain": "demographic", "age": {"min": 18}}, {}),
    (
        "inclusion",
        "經診斷為類風濕關節炎且病程達六個月以上。",
        "RA 病程≥6月",
        {"domain": "condition", "concept": "rheumatoid arthritis", "quantifier": "any", "window": {"to_days": -180}},
        {"cand": ("condition", "rheumatoid arthritis", ["類風濕關節炎"])},
    ),
    (
        "inclusion",
        "曾使用 methotrexate 治療至少六個月（連續使用，中斷不超過三十日）而療效不彰。",
        "MTX≥6月",
        {
            "domain": "medication",
            "concept": "methotrexate",
            "duration": {"min_days": 180, "gap_days": 30},
            "window": {"from_days": -730, "to_days": 0},
        },
        {"cand": ("medication", "methotrexate", ["MTX"])},
    ),
    (
        "inclusion",
        "曾使用 methotrexate 以外之另一種傳統型疾病修飾抗風濕藥物至少六個月（中斷不超過三十日）而療效不彰。",
        "另一 csDMARD≥6月",
        {
            "domain": "medication",
            "concept": "conventional synthetic dmard other than methotrexate",
            "duration": {"min_days": 180, "gap_days": 30},
            "window": {"from_days": -730, "to_days": 0},
        },
        {
            "cand": (
                "medication",
                "conventional synthetic dmard other than methotrexate",
                ["hydroxychloroquine", "sulfasalazine", "leflunomide"],
            )
        },
    ),
    (
        "inclusion",
        "申請前九十日內之疾病活動度 DAS28 大於5.1。",
        "DAS28>5.1（90日內）",
        {**DAS, "value": {"op": ">", "num": 5.1}, "window": {"from_days": -90, "to_days": 0}},
        {"action": "安排 DAS28 評估（關節數、ESR/CRP、病人整體評估）", "time_sensitive": True},
    ),
    (
        "inclusion",
        "申請前九十日至一百八十日間另有一次 DAS28 評估大於5.1（與最近一次評估間隔至少一個月）。",
        "前次 DAS28>5.1",
        {**DAS, "value": {"op": ">", "num": 5.1}, "window": {"from_days": -180, "to_days": -90}},
        {"action": "調閱前次 DAS28 評估紀錄"},
    ),
    (
        "exclusion",
        "一年內有活動性結核病。",
        "活動性結核",
        {
            "domain": "condition",
            "concept": "active tuberculosis",
            "quantifier": "any",
            "window": {"from_days": -365, "to_days": 0},
        },
        {"cand": ("condition", "active tuberculosis", ["肺結核"])},
    ),
    (
        "exclusion",
        "一年內丙型干擾素釋放試驗（IGRA）陽性而未接受潛伏結核預防性治療。",
        "IGRA陽性未治療",
        {
            "op": "and",
            "args": [
                {
                    "domain": "observation",
                    "concept": "igra",
                    "quantifier": "latest",
                    "value": {"op": "in", "codes": ["LA6576-8"]},
                    "window": {"from_days": -365, "to_days": 0},
                },
                {
                    "op": "not",
                    "args": [
                        {
                            "domain": "medication",
                            "concept": "tb preventive therapy",
                            "quantifier": "any",
                            "window": {"from_days": -365, "to_days": 0},
                        }
                    ],
                },
            ],
        },
        {
            "cand": [
                ("observation", "igra", ["丙型干擾素釋放試驗"]),
                ("medication", "tb preventive therapy", ["isoniazid"]),
            ],
            "action": "確認潛伏結核治療計畫",
        },
    ),
    (
        "exclusion",
        "B型肝炎表面抗原（HBsAg）陽性而半年內未併用抗病毒藥物預防。",
        "HBsAg陽性未預防",
        {
            "op": "and",
            "args": [
                {
                    "domain": "observation",
                    "concept": "hbsag",
                    "quantifier": "latest",
                    "value": {"op": "in", "codes": ["LA6576-8"]},
                },
                {
                    "op": "not",
                    "args": [
                        {
                            "domain": "medication",
                            "concept": "hbv antiviral prophylaxis",
                            "quantifier": "any",
                            "window": {"from_days": -180, "to_days": 0},
                        }
                    ],
                },
            ],
        },
        {
            "cand": [
                ("observation", "hbsag", ["B肝表面抗原"]),
                ("medication", "hbv antiviral prophylaxis", ["entecavir"]),
            ],
            "action": "會診肝膽科評估預防性抗病毒治療",
        },
    ),
    (
        "exclusion",
        "五年內有惡性腫瘤病史。",
        "五年內惡性腫瘤",
        {
            "domain": "condition",
            "concept": "malignancy",
            "quantifier": "any",
            "window": {"from_days": -1825, "to_days": 0},
        },
        {"cand": ("condition", "malignancy", ["惡性腫瘤", "癌症"])},
    ),
    (
        "exclusion",
        "懷孕中（二百八十日內有懷孕診斷）。",
        "懷孕",
        {
            "domain": "condition",
            "concept": "pregnancy",
            "quantifier": "any",
            "window": {"from_days": -280, "to_days": 0},
        },
        {"cand": ("condition", "pregnancy", ["懷孕"]), "time_sensitive": True},
    ),
    (
        "exclusion",
        "三十日內有嚴重感染（敗血症、肺炎）。",
        "30日內嚴重感染",
        {
            "domain": "condition",
            "concept": "serious infection",
            "quantifier": "any",
            "window": {"from_days": -30, "to_days": 0},
        },
        {"cand": ("condition", "serious infection", ["敗血症", "肺炎"]), "time_sensitive": True},
    ),
    (
        "exclusion",
        "有心衰竭（NYHA 第三或第四級）病史。",
        "心衰竭",
        {"domain": "condition", "concept": "heart failure", "quantifier": "any"},
        {"cand": ("condition", "heart failure", ["心衰竭"]), "fallback": "NYHA 分級由醫師確認"},
    ),
    (
        "renewal",
        "續用申請前九十日內須有完整之 DAS28 評估。",
        "續用 DAS28 評估",
        {**DAS, "value": {"op": ">=", "num": 0}, "window": {"from_days": -90, "to_days": 0}},
        {"action": "安排 DAS28 評估", "time_sensitive": True},
    ),
    (
        "renewal",
        "續用時療效良好：DAS28 較治療前下降至少1.2，或降至3.2以下。",
        "療效良好",
        {**DAS, "value": {"op": "<=", "num": 3.2}, "window": {"from_days": -90, "to_days": 0}},
        {
            "class": "note",
            "note_question": "病歷是否記載目前生物製劑治療反應良好（DAS28 較治療前下降至少 1.2 或降至 3.2 以下）？",
            "fallback": "結構化近似：九十日內最近一次 DAS28 ≤ 3.2",
            "action": "醫師確認療效（DAS28 變化）",
        },
    ),
    (
        "documentation",
        "一年內B型肝炎表面抗原（HBsAg）檢驗報告。",
        "HBsAg 報告",
        {"domain": "observation", "concept": "hbsag", "quantifier": "any", "window": {"from_days": -365, "to_days": 0}},
        {"cand": ("observation", "hbsag", []), "suggested_order": "14032C B型肝炎表面抗原"},
    ),
    (
        "documentation",
        "一年內B型肝炎核心抗體（anti-HBc）檢驗報告。",
        "anti-HBc 報告",
        {
            "domain": "observation",
            "concept": "anti-hbc",
            "quantifier": "any",
            "window": {"from_days": -365, "to_days": 0},
        },
        {"cand": ("observation", "anti-hbc", ["B肝核心抗體"]), "suggested_order": "14035C B型肝炎核心抗體"},
    ),
    (
        "documentation",
        "六個月內結核篩檢：丙型干擾素釋放試驗（IGRA）或胸部X光報告。",
        "結核篩檢",
        {
            "op": "or",
            "args": [
                {
                    "domain": "observation",
                    "concept": "igra",
                    "quantifier": "any",
                    "window": {"from_days": -180, "to_days": 0},
                },
                {
                    "domain": "report",
                    "concept": "chest x-ray",
                    "quantifier": "any",
                    "window": {"from_days": -180, "to_days": 0},
                },
            ],
        },
        {
            "cand": [("observation", "igra", []), ("report", "chest x-ray", ["胸部X光"])],
            "suggested_order": "12184C 丙型干擾素釋放試驗 或 32001C 胸部X光",
        },
    ),
    (
        "documentation",
        "九十日內紅血球沉降速率（ESR）或C反應蛋白（CRP）檢驗報告。",
        "ESR/CRP 報告",
        {
            "domain": "observation",
            "concept": "esr or crp",
            "quantifier": "any",
            "window": {"from_days": -90, "to_days": 0},
        },
        {"cand": ("observation", "esr or crp", ["ESR", "CRP"]), "suggested_order": "09032C CRP 或 08005C ESR"},
    ),
]
SECTIONS = {
    "inclusion": "8.2.4.1 初次申請給付條件",
    "exclusion": "8.2.4.2 不得申請情形",
    "renewal": "8.2.4.3 續用條件",
    "documentation": "8.2.4.4 申請應檢附資料",
}

# --- ONC-OSI: osimertinib for EGFR-mutated NSCLC (TWPAS cancer-drug program, phase 6) -------------------------------
ONC_TITLE = "全民健康保險藥品給付規定 9.20 非小細胞肺癌 EGFR 酪胺酸激酶抑制劑 osimertinib（合成範例）"
ONC_FRONT = [
    ("9 抗癌瘤藥物", "本文件為 TrialBox 測試用之虛構範例，依公開給付規定之結構撰寫，並非健保署正式公告內容。"),
    (
        "9.20 osimertinib",
        "限用於 EGFR 基因突變之局部晚期或轉移性非小細胞肺癌第一線治療，須經事前審查（TWPAS）核准後使用。",
    ),
]
ONC_SECTIONS = {
    "inclusion": "9.20.1 給付條件",
    "exclusion": "9.20.2 不得申請情形",
    "documentation": "9.20.3 申請應檢附資料",
}
ONC_RULES: list[tuple[str, str, str, dict[str, Any], dict[str, Any]]] = [
    ("inclusion", "病患年齡須滿十八歲。", "年齡≥18", {"domain": "demographic", "age": {"min": 18}}, {}),
    (
        "inclusion",
        "經病理或細胞學診斷為非小細胞肺癌。",
        "NSCLC 診斷",
        {"domain": "condition", "concept": "non-small cell lung cancer", "quantifier": "any"},
        {"cand": ("condition", "non-small cell lung cancer", ["非小細胞肺癌"])},
    ),
    (
        "inclusion",
        "腫瘤組織檢測證實具 EGFR 基因突變（exon 19 deletion 或 L858R）。",
        "EGFR 突變陽性",
        {
            "domain": "observation",
            "concept": "egfr mutation",
            "quantifier": "latest",
            "value": {"op": "in", "codes": ["LA9633-4"]},
        },
        {"cand": ("observation", "egfr mutation", ["EGFR基因突變"]), "action": "安排 EGFR 基因檢測"},
    ),
    (
        "inclusion",
        "ECOG 體能狀態為 0 至 1 分。",
        "ECOG 0–1",
        {"domain": "demographic", "quantifier": "any"},
        {
            "class": "note",
            "note_question": "病歷是否記載 ECOG 體能狀態為 0 或 1 分？",
            "fallback": "無結構化替代；由醫師於門診確認",
            "action": "醫師記錄 ECOG 體能狀態",
        },
    ),
    (
        "exclusion",
        "申請前一年內已使用 EGFR 酪胺酸激酶抑制劑者（本申請限第一線治療）。",
        "已用 EGFR-TKI",
        {
            "domain": "medication",
            "concept": "egfr tyrosine kinase inhibitor",
            "quantifier": "any",
            "window": {"from_days": -365, "to_days": 0},
        },
        {"cand": ("medication", "egfr tyrosine kinase inhibitor", ["osimertinib", "gefitinib", "erlotinib"])},
    ),
    (
        "documentation",
        "二個月內胸部影像檢查報告。",
        "胸部影像報告",
        {"domain": "report", "concept": "chest x-ray", "quantifier": "any", "window": {"from_days": -60, "to_days": 0}},
        {"cand": ("report", "chest x-ray", ["胸部X光"]), "suggested_order": "32001C 胸部X光 或 33070B 胸部電腦斷層"},
    ),
    (
        "documentation",
        "一年內 EGFR 基因檢測報告。",
        "EGFR 檢測報告",
        {
            "domain": "observation",
            "concept": "egfr mutation",
            "quantifier": "any",
            "window": {"from_days": -365, "to_days": 0},
        },
        {"cand": ("observation", "egfr mutation", []), "suggested_order": "12191B EGFR 基因突變檢測"},
    ),
]

SPECS: dict[str, dict[str, Any]] = {
    "RA-BIO": {
        "title": TITLE,
        "front": FRONT,
        "sections": SECTIONS,
        "rules": RULES,
        "pdf": "RA-BIO_給付規定_2026.pdf",
        "template": "RA-BIO/附表十五.docx",
    },
    "ONC-OSI": {
        "title": ONC_TITLE,
        "front": ONC_FRONT,
        "sections": ONC_SECTIONS,
        "rules": ONC_RULES,
        "pdf": "ONC-OSI_給付規定_2026.pdf",
        "template": None,
    },
}


def by_section(spec: dict[str, Any] | None = None) -> dict[str, list[str]]:
    spec = spec or SPECS["RA-BIO"]
    out: dict[str, list[str]] = {k: [] for k in spec["sections"]}
    for sec, text, *_ in spec["rules"]:
        out[sec].append(text)
    return out


def cassette(spec: dict[str, Any] | None = None) -> dict[str, Any]:
    spec = spec or SPECS["RA-BIO"]
    crit = []
    counts: dict[str, int] = {}
    for sec, text, label, logic, ex in spec["rules"]:
        counts[sec] = counts.get(sec, 0) + 1
        cands = ex.get("cand") or []
        if isinstance(cands, tuple):
            cands = [cands]
        c: dict[str, Any] = {
            "text": text,
            "label": label,
            "source_ref": f"{spec['sections'][sec]} #{counts[sec]}",
            "kind": sec,
            "class": ex.get("class", "structured"),
            "time_sensitive": bool(ex.get("time_sensitive", False)),
            "logic": logic,
            "concept_candidates": [{"domain": d, "name": n, "synonyms": s} for d, n, s in cands],
        }
        for k in ("note_question", "fallback", "action", "suggested_order"):
            if ex.get(k):
                c[k] = ex[k]
        crit.append(c)
    return {"criteria": crit}


def _fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("NotoTC", str(FONT / "NotoSansTC-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("NotoTC-Bold", str(FONT / "NotoSansTC-Bold.ttf")))
    return "NotoTC", "NotoTC-Bold"


def write_pdf(path: Path, spec: dict[str, Any] | None = None) -> None:
    spec = spec or SPECS["RA-BIO"]
    from reportlab import rl_config

    rl_config.invariant = 1
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    reg, bold = _fonts()
    body = ParagraphStyle("b", fontName=reg, fontSize=10.5, leading=15, wordWrap="CJK")
    h1 = ParagraphStyle("h1", parent=body, fontName=bold, fontSize=13, leading=18, spaceBefore=8)
    h2 = ParagraphStyle("h2", parent=body, fontName=bold, fontSize=11.5, leading=16, spaceBefore=6)
    title = ParagraphStyle("t", parent=body, fontName=bold, fontSize=15, leading=21)
    story = [Paragraph(spec["title"], title), Spacer(1, 10)]
    for head, text in spec["front"]:
        story += [Paragraph(head, h1), Paragraph(text, body)]
    for sec, items in by_section(spec).items():
        story.append(Paragraph(spec["sections"][sec], h2))
        story += [Paragraph(f"{i}. {t}", body) for i, t in enumerate(items, start=1)]
    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        title=spec["title"],
        author="TrialBox synthetic",
        creator="TrialBox",
        producer="TrialBox",
    ).build(story)


def write_docx(path: Path, spec: dict[str, Any] | None = None) -> None:
    spec = spec or SPECS["RA-BIO"]
    import docx
    from tb_common.deterministic import normalize_ooxml

    d = docx.Document()
    d.core_properties.author = "TrialBox synthetic"
    d.add_heading(spec["title"], 0)
    for head, text in spec["front"]:
        d.add_heading(head, 1)
        d.add_paragraph(text)
    for sec, items in by_section(spec).items():
        d.add_heading(spec["sections"][sec], 2)
        for i, t in enumerate(items, start=1):
            d.add_paragraph(f"{i}. {t}")
    import io

    buf = io.BytesIO()
    d.save(buf)
    path.write_bytes(normalize_ooxml(buf.getvalue()))


def write_template(path: Path) -> None:
    """附表十五 (synthetic layout) as a docxtpl template; unknown values render as highlighted [待補] (SPEC §9.4).
    Every field uses ``{{r …}}`` (RichText) and table loops use ``{%tr …%}`` rows of their own."""
    import io

    import docx
    from docx.shared import Pt
    from tb_common.deterministic import normalize_ooxml

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Noto Sans TC"
    st.font.size = Pt(10.5)
    d.add_heading("附表十五 類風濕關節炎病患使用生物製劑申請表（合成範例）", 1)
    d.add_paragraph("申請類別：{{r apply_type }}　　申請日期：{{r run_date }}　　醫院：{{r site_name }}")
    d.add_heading("一、病患資料（院內使用）", 2)
    t = d.add_table(rows=3, cols=4)
    t.style = "Table Grid"
    cells = [
        ("病歷號", "{{r mrn }}", "性別", "{{r sex }}"),
        ("出生日期", "{{r birth_date }}", "年齡", "{{r age }}"),
        ("主治醫師", "{{r physician }}", "科別", "{{r department }}"),
    ]
    for r, row in enumerate(cells):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_heading("二、診斷", 2)
    d.add_paragraph("診斷：{{r diagnosis }}（ICD-10-CM {{r diagnosis_code }}），首次診斷日期：{{r diagnosis_date }}")

    def loop_table(headers: tuple[str, ...], var: str, item: str, keys: tuple[str, ...]) -> None:
        tb = d.add_table(rows=4, cols=len(headers))
        tb.style = "Table Grid"
        for c, h in enumerate(headers):
            tb.cell(0, c).text = h
        tb.cell(1, 0).text = f"{{%tr for {item} in {var} %}}"
        for c, k in enumerate(keys):
            tb.cell(2, c).text = f"{{{{r {item}.{k} }}}}"
        tb.cell(3, 0).text = "{%tr endfor %}"

    d.add_heading("三、傳統型疾病修飾抗風濕藥物（DMARDs）使用紀錄", 2)
    loop_table(("藥品", "劑量", "起日", "迄日", "使用日數"), "dmards", "m", ("name", "dose", "start", "end", "days"))
    d.add_paragraph("類固醇（prednisolone）：{{r prednisolone }}")
    d.add_heading("四、疾病活動度 DAS28", 2)
    loop_table(
        ("評估", "日期", "TJC28", "SJC28", "ESR/CRP", "PtGA", "DAS28", "方法"),
        "das28",
        "a",
        ("label", "date", "tjc", "sjc", "marker", "ptga", "score", "method"),
    )
    d.add_heading("五、感染篩檢", 2)
    d.add_paragraph("HBsAg：{{r hbsag }}；anti-HBc：{{r ahbc }}；結核篩檢：{{r tb_screen }}")
    d.add_heading("六、不得申請情形檢核", 2)
    d.add_paragraph("{%p for x in exclusions %}")
    d.add_paragraph("☐/☑ {{r x.label }}：{{r x.status }}")
    d.add_paragraph("{%p endfor %}")
    d.add_heading("七、臨床病程摘要", 2)
    d.add_paragraph("{{r course }}")
    d.add_paragraph("本申請書由 TrialBox 依病歷資料預填，標示 [待補] 之欄位須由醫師補齊並親自簽章；系統不代為簽署。")
    d.add_paragraph("醫師簽章：＿＿＿＿＿＿＿＿　　日期：＿＿＿＿＿＿")
    d.add_paragraph("{{r footer }}")
    buf = io.BytesIO()
    d.save(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(normalize_ooxml(buf.getvalue()))


def check(pdf: Path, spec: dict[str, Any] | None = None) -> int:
    spec = spec or SPECS["RA-BIO"]
    from doc_parser.parser import parse
    from llm_stub.handlers import norm

    parsed = parse(pdf.read_bytes(), pdf.name)
    got = parsed.ie_block
    bad = 0
    for sec, want in by_section(spec).items():
        have = getattr(got, sec) or []
        if [norm(x) for x in have] != [norm(x) for x in want]:
            bad += 1
            print(f"{sec}: parsed {have!r}\n  want {want!r}")
    n = len(spec["sections"])
    print(f"{pdf.name}: language={parsed.language}; sections ok: {n - bad}/{n}")
    return 1 if bad else 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args(argv)
    rc = 0
    for rid, spec in SPECS.items():
        pdf = ROOT / "tests/fixtures/protocols" / spec["pdf"]
        if not args.check:
            write_pdf(pdf, spec)
            write_docx(pdf.with_suffix(".docx"), spec)
            (ROOT / f"services/llm_stub/cassettes/ir_extract/{rid}.json").write_text(
                json.dumps(cassette(spec), ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
            )
            if spec["template"]:
                write_template(ROOT / "services/orchestrator/reports/templates" / spec["template"])
        rc |= check(pdf, spec)
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
