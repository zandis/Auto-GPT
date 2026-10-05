"""Write the synthetic source documents the compiler is tested on (fictional; modelled on public trial designs).

* ``GZQO_protocol_v3.pdf`` / ``.docx`` — retatrutide-in-gout phase 3 protocol excerpt (English), §5 eligibility.
* ``GZQO_questions.xlsx`` — a sponsor feasibility question list (FEAS input variant).
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CASSETTE = ROOT / "services" / "llm_stub" / "cassettes" / "ir_extract"

FRONT = [
    (
        "1 Protocol Summary",
        "Protocol J1I-MC-GZQO (synthetic excerpt). A Phase 3, randomized, double-blind, "
        "placebo-controlled study of once-weekly retatrutide in participants with gout and obesity or overweight. "
        "This document is a fictional excerpt produced for TrialBox testing; it is not the sponsor's protocol.",
    ),
    (
        "2 Objectives and Endpoints",
        "Primary: proportion of participants with serum urate below 6.0 mg/dL at week "
        "52. Key secondary: gout flare rate, change in body weight.",
    ),
    (
        "3 Study Design",
        "Approximately 900 participants will be randomized 1:1:1 to retatrutide 4 mg, 8 mg or placebo "
        "for 52 weeks after a 4-week screening period.",
    ),
]


def criteria(ruleset: str) -> tuple[list[str], list[str]]:
    data = json.loads((CASSETTE / f"{ruleset}.json").read_text(encoding="utf-8"))["criteria"]
    inc = [c["text"] for c in data if c["kind"] == "inclusion"]
    exc = [c["text"] for c in data if c["kind"] == "exclusion"]
    return inc, exc


def write_pdf(path: Path, title: str, inc: list[str], exc: list[str]) -> None:
    from reportlab import rl_config

    rl_config.invariant = 1  # byte-reproducible PDF (no timestamps / random ids)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for head, body in FRONT:
        story += [Paragraph(head, styles["Heading1"]), Paragraph(body, styles["BodyText"]), Spacer(1, 8)]
    story.append(Paragraph("5 Study Population", styles["Heading1"]))
    for head, items in (("5.1 Inclusion Criteria", inc), ("5.2 Exclusion Criteria", exc)):
        story.append(Paragraph(head, styles["Heading2"]))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(t, styles["BodyText"])) for t in items],  # type: ignore[misc]
                bulletType="1",
                bulletFormat="%s.",
            )
        )
    story += [
        Paragraph("6 Study Intervention", styles["Heading1"]),
        Paragraph("Retatrutide is administered subcutaneously once weekly.", styles["BodyText"]),
    ]
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        title=title,
        author="TrialBox synthetic",
        subject="test",
        creator="TrialBox",
        producer="TrialBox",
    )
    doc.build(story)


def write_docx(path: Path, title: str, inc: list[str], exc: list[str]) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "TrialBox synthetic"
    d.add_heading(title, 0)
    for head, body in FRONT:
        d.add_heading(head, 1)
        d.add_paragraph(body)
    d.add_heading("5 Study Population", 1)
    for head, items in (("5.1 Inclusion Criteria", inc), ("5.2 Exclusion Criteria", exc)):
        d.add_heading(head, 2)
        for i, t in enumerate(items, start=1):
            d.add_paragraph(f"{i}. {t}")
    d.save(str(path))


def write_question_xlsx(path: Path, inc: list[str], exc: list[str]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Feasibility"
    ws.append(["#", "Section", "Question / criterion", "Site answer"])
    n = 1
    for sec, items in (("Inclusion Criteria", inc), ("Exclusion Criteria", exc)):
        for t in items:
            ws.append([n, sec, t, ""])
            n += 1
    wb.save(str(path))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=ROOT / "tests" / "fixtures" / "protocols")
    args = ap.parse_args(argv)
    args.out.mkdir(parents=True, exist_ok=True)
    inc, exc = criteria("GZQO")
    title = "J1I-MC-GZQO Protocol v3.0 (synthetic excerpt)"
    write_pdf(args.out / "GZQO_protocol_v3.pdf", title, inc, exc)
    write_docx(args.out / "GZQO_protocol_v3.docx", title, inc, exc)
    write_question_xlsx(args.out / "GZQO_questions.xlsx", inc, exc)
    print(f"wrote protocol documents to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
