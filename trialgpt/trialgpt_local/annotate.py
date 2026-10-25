"""Blinded physician annotation workbooks (shared by the trialgpt.app head-to-head and the CGRD validation).

``write_workbook`` produces an XLSX with one row per item to judge (patient summary, trial, criterion when given) and
drop-down answer columns; which system proposed an item never appears in the workbook — it stays in a separate key
(JSON) that the scorer joins after adjudication. Rows are shuffled with a fixed seed so the order reveals nothing.
"""

from __future__ import annotations

import io
import json
import random
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TRIAL_CHOICES = ("2 eligible", "1 potentially eligible / needs data", "0 not eligible or not relevant")
CRITERION_CHOICES = {
    "inclusion": ("included", "not included", "not enough information", "not applicable"),
    "exclusion": ("excluded", "not excluded", "not enough information", "not applicable"),
}


def write_workbook(items: list[dict[str, Any]], path: Path, *, title: str, seed: int = 29) -> Path:
    """``items``: dicts with ``item_id``, ``patient_id``, ``patient``, ``nct_id``, ``trial_title``, ``trial_text`` and
    optionally ``kind`` + ``criterion`` (criterion-level rows). Writes ``path`` (XLSX) and ``path``.key.json."""
    rows = list(items)
    random.Random(seed).shuffle(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "judge"
    crit = any(r.get("criterion") for r in rows)
    head = ["row", "patient", "patient summary", "trial", "trial (title, summary, criteria)"]
    head += ["criterion type", "criterion", "verdict"] if crit else ["eligibility (0/1/2)"]
    head += ["comment", "minutes"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE6F7")
    for i, r in enumerate(rows, start=1):
        line: list[Any] = [i, r["patient_id"], r["patient"], r["nct_id"], f"{r['trial_title']}\n\n{r['trial_text']}"]
        answer: list[Any] = [r.get("kind", ""), r.get("criterion", ""), ""] if crit else [""]
        ws.append([*line, *answer, "", ""])
    widths = [6, 14, 60, 13, 70] + ([11, 50, 22] if crit else [26]) + [30, 9]
    for col, w in zip("ABCDEFGHIJ", widths, strict=False):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    answer_col = "H" if crit else "F"
    if crit:
        for kind, choices in CRITERION_CHOICES.items():
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            for i, r in enumerate(rows, start=2):
                if r.get("kind") == kind:
                    dv.add(f"{answer_col}{i}")
    else:
        dv = DataValidation(type="list", formula1='"' + ",".join(TRIAL_CHOICES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{answer_col}2:{answer_col}{len(rows) + 1}")
    ws.freeze_panes = "C2"
    info = wb.create_sheet("instructions")
    for text in (
        title,
        "",
        "Judge each row from the patient summary and the trial text only. Which system proposed a row is hidden.",
        "Trial level: 2 = eligible (meets inclusion, no exclusion); 1 = potentially eligible / missing data;",
        "0 = not eligible or not relevant to this patient.",
        "Criterion level: pick the verdict for this criterion as the TrialGPT paper defines it.",
        "Fill 'minutes' with the time the row took (optional). Do not reorder or delete rows.",
    ):
        info.append([text])
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(buf.getvalue())
    key = [
        {"row": i, **{k: v for k, v in r.items() if k not in ("patient", "trial_text")}} for i, r in enumerate(rows, 1)
    ]
    Path(str(path) + ".key.json").write_text(json.dumps(key, indent=1, ensure_ascii=False), encoding="utf-8")
    return path


def read_workbook(path: Path) -> list[dict[str, Any]]:
    """Answers joined with the key: rows with an empty verdict are kept with ``answer=None``."""
    key = {k["row"]: k for k in json.loads(Path(str(path) + ".key.json").read_text(encoding="utf-8"))}
    ws = load_workbook(path, data_only=True)["judge"]
    head = [c.value for c in ws[1]]
    ans_col = head.index("verdict") if "verdict" in head else head.index("eligibility (0/1/2)")
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        raw = row[ans_col]
        answer: Any = None
        if raw not in (None, ""):
            s = str(raw).strip().lower()
            answer = int(s[0]) if s[:1] in ("0", "1", "2") and "verdict" not in head else s
        minutes = row[head.index("minutes")]
        out.append(
            {
                **key[int(row[0])],
                "answer": answer,
                "comment": row[head.index("comment")],
                "minutes": float(minutes) if isinstance(minutes, (int, float)) else None,
            }
        )
    return out
