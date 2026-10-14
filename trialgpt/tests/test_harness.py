"""Benchmark, pipeline, web API, trialgpt.app head-to-head, CGRD validation — all offline with the mock LLM."""

from __future__ import annotations

import hashlib
import json
import sqlite3
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import trial
from trialgpt_local import annotate, benchmark, ctgov, datasets, pipeline, retrieval


@pytest.fixture(autouse=True)
def fake_medcpt(monkeypatch: pytest.MonkeyPatch) -> None:
    """Deterministic bag-of-words vectors instead of MedCPT (no model download in tests)."""

    def vec(text: str) -> np.ndarray:
        v = np.zeros(768, dtype=np.float32)
        for w in text.lower().split():
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 768] += 1
        return v

    monkeypatch.setattr(
        retrieval,
        "encode_articles",
        lambda docs, batch=16: (
            np.stack([vec(f"{d.get('title')} {d.get('text')}") for d in docs]) if docs else np.zeros((0, 768))
        ),
    )
    monkeypatch.setattr(retrieval, "encode_queries", lambda texts: np.stack([vec(t) for t in texts]))
    pipeline._INDEXES.clear()


def test_benchmark_rank_and_report(mini_cohort: str) -> None:
    d = benchmark.run_rank(mini_cohort, None, None)
    matching = json.loads(next(d.glob("matching_results_*.json")).read_text())
    assert set(matching) == {"sigir-1", "sigir-2"} and "NCT1" in matching["sigir-1"]["2"]  # NCBI layout
    ev = benchmark.evaluate(mini_cohort, "mock-model")
    combined = ev["metrics"]["combined"]
    # sigir-1: NCT1 (label 2) has no failed criterion -> ranked above NCT2 (label 1, one NEGATIVE) -> AUROC 1
    assert combined["auroc"] == 1.0 and combined["ndcg@10"] is not None
    assert ev["diagnostics"]["criterion_parse_error_rate"] > 0  # the GARBLED criterion is counted, not hidden
    text = benchmark.report("mock-model")
    assert "Paper, GPT-4" in text and "0.7275" in text
    # resumable: a second run makes no new LLM call
    from tests.conftest import CALLS

    n = len(CALLS)
    benchmark.run_rank(mini_cohort, None, None)
    assert len(CALLS) == n


def test_slice_keeps_every_label(mini_cohort: str) -> None:
    co = datasets.load(mini_cohort)
    sel = benchmark.select(co, 1, 2)
    ((pid, ts),) = sel.items()
    labels = {co.labels_for(pid)[t["NCTID"]] for t in ts}
    assert len(ts) == 2 and len(labels) == 2


def test_pipeline_on_cohort_corpus(mini_cohort: str) -> None:
    from trialgpt_local.llm import LLMClient

    run = pipeline.run_patient("A man with gout and high serum urate.", client=LLMClient(), source=mini_cohort, top_k=3)
    assert run.candidates == 4 and len(run.ranked) == 3
    assert run.ranked == sorted(run.ranked, key=lambda m: -m.score)
    assert run.to_json()["ranked"][0]["score"] == run.ranked[0].score


def test_live_pipeline_prefers_country(monkeypatch: pytest.MonkeyPatch) -> None:
    from trialgpt_local.llm import LLMClient

    pool = [
        trial("NCT10", ["gout"], ["x"], countries=["France"]),
        trial("NCT11", ["gout"], ["x"], countries=["Taiwan"]),
    ]
    monkeypatch.setattr(ctgov, "search", lambda conds, **kw: pool)
    run = pipeline.run_patient(
        "gout patient note here", client=LLMClient(), source="live", top_k=2, prefer_country="Taiwan"
    )
    assert run.retrieved[0] == "NCT11"


def test_live_search_falls_back_to_full_text() -> None:
    import httpx

    from tests.conftest import REAL_CLIENT

    seen: list[tuple[str, str]] = []

    def study(n: int) -> dict[str, Any]:
        return {
            "protocolSection": {
                "identificationModule": {"nctId": f"NCT{n:08d}", "briefTitle": f"T{n}"},
                "eligibilityModule": {"eligibilityCriteria": "Inclusion Criteria:\n\n* adult"},
            }
        }

    def registry(request: httpx.Request) -> httpx.Response:
        field = "query.cond" if "query.cond" in request.url.params else "query.term"
        seen.append((field, request.url.params[field]))
        n = 3 if field == "query.cond" and request.url.params[field] == "progressive dysphagia" else 20
        base = 100 if field == "query.term" else 0
        return httpx.Response(200, json={"studies": [study(base + i) for i in range(n)]})

    pool = ctgov.search(
        ["dysphagia", "progressive dysphagia"], http=REAL_CLIENT(transport=httpx.MockTransport(registry))
    )
    assert seen == [
        ("query.cond", "dysphagia"),
        ("query.cond", "progressive dysphagia"),
        ("query.term", "progressive dysphagia"),
    ]
    assert len(pool) == 20 + 20  # 0..19 from the condition searches, 100..119 from full text, de-duplicated


def test_split_criteria_free_form_headings() -> None:
    text = (
        "Patients are eligible to be included in the study only if all the following criteria apply:\n\n"
        "* Adults \\>18 years\n* Gout\n  diagnosed by a rheumatologist\n\n"
        "Patients are excluded from the study if any of the following criteria apply:\n\n* Pregnant\n* Dialysis"
    )
    inc, exc = ctgov.split_criteria(text)
    assert inc.split("\n\n") == ["Adults >18 years", "Gout diagnosed by a rheumatologist"]
    assert exc.split("\n\n") == ["Pregnant", "Dialysis"]
    inc2, exc2 = ctgov.split_criteria("Inclusion Criteria:\n\n1. Age 20-75\n2. RA\n\nKey Exclusion Criteria:\n\n- TB")
    assert inc2.split("\n\n") == ["Age 20-75", "RA"] and exc2 == "TB"


def test_blinded_workbook_round_trip(tmp_path: Path) -> None:
    items = [
        {
            "item_id": f"p|N{i}",
            "patient_id": "p",
            "patient": "note",
            "nct_id": f"N{i}",
            "trial_title": "t",
            "trial_text": "x",
            "systems": ["local"] if i % 2 else ["trialgpt.app"],
        }
        for i in range(4)
    ]
    path = annotate.write_workbook(items, tmp_path / "w.xlsx", title="t")
    ws = load_workbook(path)["judge"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "trialgpt.app" not in flat and "local" not in flat  # system identity is only in the key
    ws.cell(2, 6, "2 eligible")
    ws.cell(3, 6, "0 not eligible or not relevant")
    ws.parent.save(path)
    got = {r["nct_id"]: r["answer"] for r in annotate.read_workbook(path)}
    assert sorted(v for v in got.values() if v is not None) == [0, 2] and list(got.values()).count(None) == 2


def test_h2h_report(mini_cohort: str, isolated: Path) -> None:
    from trialgpt_local.h2h import compare

    cap = isolated / "results" / "h2h" / "trialgpt_app"
    cap.mkdir(parents=True)
    (cap / "sigir-1.json").write_text(
        json.dumps(
            {
                "patient_id": "sigir-1",
                "ranked": ["NCT2", "NCT9"],
                "search_seconds": 7.0,
                "select_policy": "first",
                "country": "United States",
            }
        )
    )
    loc = compare.local_dir()
    (loc / "sigir-1.json").write_text(json.dumps({"ranked": [{"nct_id": "NCT1"}, {"nct_id": "NCT8"}]}))
    wb = compare.h2h_dir() / "pooled.xlsx"
    items = [
        {
            "item_id": f"sigir-1|{n}",
            "patient_id": "sigir-1",
            "patient": "x",
            "nct_id": n,
            "trial_title": "t",
            "trial_text": "x",
        }
        for n in ("NCT1", "NCT8", "NCT2", "NCT9")
    ]
    annotate.write_workbook(items, wb, title="t")
    ws = load_workbook(wb)["judge"]
    for row in range(2, 6):
        nct = ws.cell(row, 4).value
        ws.cell(row, 6, "2 eligible" if nct == "NCT1" else "0 not eligible or not relevant")
    ws.parent.save(wb)
    text = compare.report(wb)
    assert "labelled share of top-10" in text and "| local | 0.500 |" in text  # NCT1 labelled, NCT8 not
    assert "local − trialgpt.app: mean +0.100" in text  # graded P@10: 2/20 vs 0


def test_app_api(mini_cohort: str, monkeypatch: pytest.MonkeyPatch) -> None:
    from trialgpt_local import app as appmod

    c = TestClient(appmod.app)
    assert c.get("/").status_code == 200
    cfg = c.get("/api/config").json()
    assert cfg["model"] == "mock-model" and "live" in cfg["sources"] and "sigir" in cfg["sources"]
    up = c.post("/api/upload", files={"file": ("n.txt", b"A 70-year-old with gout.")}).json()
    assert up["text"].startswith("A 70")
    jid = c.post(
        "/api/match", json={"note": "A 60-year-old man with gout and urate 9.", "source": "sigir", "top_k": 2}
    ).json()["id"]
    import time

    for _ in range(100):
        j = c.get(f"/api/jobs/{jid}").json()
        if j["state"] != "running":
            break
        time.sleep(0.05)
    assert j["state"] == "done", j["error"]
    assert len(j["result"]["ranked"]) == 2 and j["log"][-1]["msg"] == "Done"
    assert c.get("/static/../app.py").status_code == 404
    assert c.post("/api/match", json={"note": "short"}).status_code == 422


def _cgrd_db(path: Path) -> None:
    con = sqlite3.connect(path)
    con.executescript("""
      CREATE TABLE patient (mrn TEXT, name TEXT, id_no TEXT, birth_date TEXT, sex TEXT, death_date TEXT, phone TEXT);
      CREATE TABLE diagnosis (diag_no TEXT, enc_no TEXT, mrn TEXT, icd10 TEXT, diag_date TEXT, recorded_date TEXT, status TEXT);
      CREATE TABLE lab (lab_no TEXT, mrn TEXT, local_code TEXT, item_name TEXT, result_value TEXT, unit TEXT, sample_datetime TEXT);
      CREATE TABLE vital (vital_no TEXT, mrn TEXT, measure_datetime TEXT, height_cm TEXT, weight_kg TEXT, sbp TEXT, dbp TEXT);
      CREATE TABLE medication (rx_no TEXT, mrn TEXT, nhi_drug_code TEXT, atc TEXT, drug_name TEXT, order_date TEXT, start_date TEXT,
                               end_date TEXT, daily_dose TEXT, dose_unit TEXT, status TEXT, dept_code TEXT);
      CREATE TABLE procedure (proc_no TEXT, mrn TEXT, order_code TEXT, proc_name TEXT, proc_date TEXT);
      CREATE TABLE report (rpt_no TEXT, mrn TEXT, exam_code TEXT, exam_name TEXT, exam_date TEXT, conclusion TEXT);
      CREATE TABLE note (note_no TEXT, mrn TEXT, enc_no TEXT, note_datetime TEXT, note_type TEXT, text TEXT);
      INSERT INTO patient VALUES ('C001','王小明','A123456789','1960-05-01','M','','0912345678'),
                                 ('C002','林小華','B223456789','1980-01-01','F','','0922345678');
      INSERT INTO diagnosis VALUES ('d1','e1','C001','M10.9','2020-01-01','','active'), ('d2','e2','C002','M06.9','2019-01-01','','active');
      INSERT INTO lab VALUES ('l1','C001','UA','Uric acid','9.1','mg/dL','2026-09-01 08:00:00');
      INSERT INTO vital VALUES ('v1','C001','2026-09-01 08:00:00','170','80','130','80');
      INSERT INTO medication VALUES ('m1','C001','','M04AA01','Allopurinol 100mg','2026-07-01','2026-07-01','2026-07-28','100','mg','active',''),
                                    ('m2','C001','','M04AA01','Allopurinol 100mg','2026-07-29','2026-07-29','2026-10-01','100','mg','active','');
      INSERT INTO note VALUES ('n1','C001','e1','2026-09-01 09:00:00','progress','右側第一蹠趾關節紅腫。持續服藥。');
    """)
    con.commit()
    con.close()


def test_cgrd_notes_and_validation_loop(isolated: Path) -> None:
    from datetime import date

    from trialgpt_local.cgrd import validate
    from trialgpt_local.cgrd.reader import Source, compose_notes, load_mapping, pseudonym

    db = isolated / "cgrd.sqlite"
    _cgrd_db(db)
    notes = compose_notes(Source(str(db), load_mapping()), ["C001", "C002"], date(2026, 10, 5))
    n = notes["C001"]
    assert n.startswith("A 66-year-old male patient.") and "Idiopathic gout" not in n and "M10.9" in n
    assert "Allopurinol 100mg 100 mg (since 2026-07-01)" in n  # refills merged into one current course
    assert "王小明" not in n and "A123456789" not in n and "1960-05-01" not in n and "0912345678" not in n
    assert "右側第一蹠趾關節紅腫。" in n
    assert pseudonym("C001") == pseudonym("C001") != pseudonym("C002")
    trials = isolated / "trials.json"
    trials.write_text(
        json.dumps(
            [
                trial(
                    "NCTG",
                    ["Adults with gout REQUIRES:M10.9", "Uric acid above 8"],
                    ["Pregnancy"],
                    diseases_list=["gout"],
                )
            ]
        )
    )
    validate.select(
        str(db), None, str(trials), "t1", per_trial=1, controls=1, index_date=date(2026, 10, 5), mrns_file=None
    )
    d = validate.run_dir("t1")
    pairs = json.loads((d / "pairs.json").read_text())
    assert [p["pool"] for p in pairs] == ["likely", "control"] and pairs[0]["patient_id"] == pseudonym("C001")
    assert oct((d / "notes.json").stat().st_mode & 0o777) == "0o600"
    validate.run_matching("t1")
    before = list((isolated / "cache").rglob("*.json"))
    assert before == []  # CGRD prompts never reach the disk cache
    validate.workbooks("t1", criteria_pairs=2)
    ws = load_workbook(d / "trial.xlsx")["judge"]
    for row in range(2, ws.max_row + 1):
        pid = ws.cell(row, 2).value
        ws.cell(row, 6, "2 eligible" if pid == pseudonym("C001") else "0 not eligible or not relevant")
    ws.parent.save(d / "trial.xlsx")
    cw = load_workbook(d / "criteria.xlsx")["judge"]
    for row in range(2, cw.max_row + 1):
        gout = cw.cell(row, 2).value == pseudonym("C001")
        kind, crit = cw.cell(row, 6).value, str(cw.cell(row, 7).value)
        verdict = (
            "not excluded"
            if kind == "exclusion"
            else ("included" if gout or "REQUIRES" not in crit else "not included")
        )
        cw.cell(row, 8, verdict)
    cw.parent.save(d / "criteria.xlsx")
    text = validate.score("t1")
    assert "AUROC (eligible vs not): 1.000" in text and "accuracy 1.000" in text


def _unused(_: Any) -> None:  # keep imports used by type checkers in sync
    return None
